"""
Axioms to Excel Converter
Intelligently extracts and organizes axiom information from markdown files into Excel
"""

import os
import re
from pathlib import Path
import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

def extract_yaml_frontmatter(content):
    """Extract YAML frontmatter from markdown content"""
    pattern = r'^---\s*\n(.*?)\n---\s*\n'
    match = re.match(pattern, content, re.DOTALL)
    if match:
        try:
            return yaml.safe_load(match.group(1))
        except yaml.YAMLError:
            return {}
    return {}

def extract_section(content, section_name):
    """Extract content from a specific markdown section"""
    pattern = rf'^##\s+{re.escape(section_name)}\s*\n(.*?)(?=^##\s|\Z)'
    match = re.search(pattern, content, re.MULTILINE | re.DOTALL)
    if match:
        return match.group(1).strip()
    return ""

def extract_formal_statement(content):
    """Extract the formal statement"""
    section = extract_section(content, "Formal Statement")
    if section:
        lines = section.split('\n')
        for line in lines:
            line = line.strip()
            if line and not line.startswith('-') and not line.startswith('*'):
                return line
    return ""

def extract_classification_emoji(classification):
    """Extract emoji from classification string"""
    if not classification:
        return ""
    emoji_pattern = r'[\U0001F300-\U0001F9FF]|[\u2600-\u26FF]|[\u2700-\u27BF]'
    match = re.search(emoji_pattern, str(classification))
    return match.group(0) if match else ""

def clean_text(text):
    """Clean text by removing wiki links and extra whitespace"""
    if not text:
        return ""
    text = re.sub(r'\[\[.*?\|(.*?)\]\]', r'\1', text)
    text = re.sub(r'\[\[(.*?)\]\]', r'\1', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def extract_dependencies(depends_on):
    """Extract dependency list"""
    if not depends_on:
        return ""
    if isinstance(depends_on, list):
        return ", ".join(depends_on)
    return str(depends_on)

def extract_enables(content):
    """Extract what this axiom enables"""
    section = extract_section(content, "Enables")
    if section:
        enables_list = []
        for line in section.split('\n'):
            line = line.strip()
            if line.startswith('-'):
                match = re.search(r'\[\[.*?\|(.*?)\]\]|\[\[(.*?)\]\]', line)
                if match:
                    enables_list.append(match.group(1) or match.group(2))
        return ", ".join(enables_list) if enables_list else clean_text(section)
    return ""

def extract_domains(domain):
    """Extract domain list"""
    if not domain:
        return ""
    if isinstance(domain, list):
        return ", ".join(domain)
    return str(domain)

def process_axiom_file(filepath):
    """Process a single axiom markdown file"""
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        
        metadata = extract_yaml_frontmatter(content)
        
        axiom_data = {
            'File': os.path.basename(filepath),
            'Axiom ID': metadata.get('axiom_id', ''),
            'Chain Position': metadata.get('chain_position', ''),
            'Classification': str(metadata.get('classification', '')).replace('\U0001F4D0', '📐').replace('\U0001F537', '🔷').replace('🟢', '🟢'),
            'Status': metadata.get('status', ''),
            'Stage': metadata.get('stage', ''),
            'Tier': metadata.get('tier', ''),
            'Domain': extract_domains(metadata.get('domain', [])),
            'Depends On': extract_dependencies(metadata.get('depends_on', [])),
            'Enables': extract_enables(content),
            'Formal Statement': clean_text(extract_formal_statement(content)),
            'UUID': metadata.get('uuid', ''),
            'Paper Refs': str(metadata.get('paper_refs', [])),
            'Collapse Radius': metadata.get('collapse_radius', ''),
        }
        
        return axiom_data
    except Exception as e:
        print(f"Error processing {filepath}: {e}")
        return None

def create_styled_excel(df, output_path):
    """Create a beautifully styled Excel file"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Axioms Overview"
    
    # Define styles
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Write headers
    for col_num, column_title in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = column_title
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    
    # Write data
    for row_num, row_data in enumerate(df.values, 2):
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num)
            # Convert value to string if it's a list or other non-primitive type
            if isinstance(value, (list, dict)):
                cell.value = str(value)
            elif pd.isna(value):
                cell.value = ""
            else:
                cell.value = value
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            
            # Color code by classification
            if col_num == 4:  # Classification column
                if 'Primitive' in str(value):
                    cell.fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
                elif 'Definition' in str(value):
                    cell.fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
                elif 'Theorem' in str(value):
                    cell.fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
    
    # Adjust column widths
    column_widths = {
        'A': 35,  # File
        'B': 12,  # Axiom ID
        'C': 12,  # Chain Position
        'D': 20,  # Classification
        'E': 15,  # Status
        'F': 10,  # Stage
        'G': 10,  # Tier
        'H': 25,  # Domain
        'I': 20,  # Depends On
        'J': 30,  # Enables
        'K': 60,  # Formal Statement
        'L': 38,  # UUID
        'M': 20,  # Paper Refs
        'N': 15,  # Collapse Radius
    }
    
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width
    
    # Freeze header row
    ws.freeze_panes = 'A2'
    
    # Add auto-filter
    ws.auto_filter.ref = ws.dimensions
    
    wb.save(output_path)
    print(f"✓ Excel file created: {output_path}")

def main():
    """Main execution function"""
    axioms_folder = r"O:\Theophysics_Master\TMSUB\GO FOLDER\_AXIOMS_001-188"
    output_file = r"O:\Theophysics_Backend\Python_Backend\Axioms_Database.xlsx"
    
    print("=" * 70)
    print("AXIOMS TO EXCEL CONVERTER")
    print("=" * 70)
    print(f"\nSource folder: {axioms_folder}")
    print(f"Output file: {output_file}\n")
    
    # Get only the 188 numbered axiom files (001-188)
    md_files = []
    for f in Path(axioms_folder).glob("*.md"):
        if f.is_file():
            # Only include files that start with a number (001-188)
            filename = f.name
            if re.match(r'^\d{3}_', filename):
                md_files.append(f)
    
    md_files = sorted(md_files)
    print(f"Found {len(md_files)} numbered axiom files (excluding case files and other documents)\n")
    
    # Process each file
    axioms_data = []
    for idx, filepath in enumerate(md_files, 1):
        print(f"Processing [{idx}/{len(md_files)}]: {filepath.name}")
        data = process_axiom_file(filepath)
        if data:
            axioms_data.append(data)
    
    # Create DataFrame
    df = pd.DataFrame(axioms_data)
    
    # Convert Chain Position to numeric for proper sorting
    df['Chain Position'] = pd.to_numeric(df['Chain Position'], errors='coerce')
    
    # Sort by chain position
    df = df.sort_values('Chain Position', na_position='last')
    
    print(f"\n✓ Processed {len(axioms_data)} axioms successfully")
    print(f"\nCreating Excel file...")
    
    # Create styled Excel
    create_styled_excel(df, output_file)
    
    # Print summary statistics
    print("\n" + "=" * 70)
    print("SUMMARY STATISTICS")
    print("=" * 70)
    print(f"Total Axioms: {len(df)}")
    print(f"\nBy Status:")
    print(df['Status'].value_counts().to_string())
    print(f"\nBy Stage:")
    print(df['Stage'].value_counts().sort_index().to_string())
    print(f"\nBy Domain (top 5):")
    domain_counts = df['Domain'].str.split(', ').explode().value_counts().head(5)
    print(domain_counts.to_string())
    print("\n" + "=" * 70)
    print("✓ COMPLETE!")
    print("=" * 70)

if __name__ == "__main__":
    main()
