"""
Comprehensive Axiom System Extractor
Extracts ALL content from axioms folder including:
- 188 numbered axioms
- Prosecution cases
- Working papers
- Case files
- All classifications and dependencies
"""

import os
import re
from pathlib import Path
import yaml
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import defaultdict

def extract_yaml_frontmatter(content):
    """Extract YAML frontmatter from markdown content"""
    pattern = r'^---\s*\n(.*?)\n---\s*\n'
    match = re.match(pattern, content, re.DOTALL)
    if match:
        try:
            return yaml.safe_load(match.group(1))
        except yaml.YAMLError:
            return {}
    return {}

def extract_section(content, section_name):
    """Extract content from a specific markdown section"""
    pattern = rf'^##\s+{re.escape(section_name)}\s*\n(.*?)(?=^##\s|\Z)'
    match = re.search(pattern, content, re.MULTILINE | re.DOTALL)
    if match:
        return match.group(1).strip()
    return ""

def clean_text(text):
    """Clean text by removing wiki links and extra whitespace"""
    if not text:
        return ""
    text = re.sub(r'\[\[.*?\|(.*?)\]\]', r'\1', text)
    text = re.sub(r'\[\[(.*?)\]\]', r'\1', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

class ComprehensiveAxiomExtractor:
    def __init__(self, axioms_folder):
        self.axioms_folder = axioms_folder
        self.numbered_axioms = []
        self.prosecution_files = []
        self.working_papers = []
        self.case_files = []
        self.other_files = []
        
    def scan_all_files(self):
        """Scan entire axioms folder for all file types"""
        print(f"Scanning folder: {self.axioms_folder}\n")
        
        all_files = list(Path(self.axioms_folder).glob("**/*.md"))
        
        for filepath in all_files:
            filename = filepath.name
            relative_path = str(filepath.relative_to(self.axioms_folder))
            
            # Categorize files
            if re.match(r'^\d{3}_', filename):
                # Numbered axioms (001-188)
                self.numbered_axioms.append(filepath)
            elif 'PROSECUT' in filename.upper() or '_PROSECUTED' in str(filepath):
                # Prosecution files
                self.prosecution_files.append(filepath)
            elif 'WORKING' in str(filepath).upper() or '_WORKING_PAPERS' in str(filepath):
                # Working papers
                self.working_papers.append(filepath)
            elif 'CASE' in str(filepath).upper() or '_CASE_FILES' in str(filepath):
                # Case files
                self.case_files.append(filepath)
            else:
                # Other markdown files
                self.other_files.append(filepath)
        
        print(f"File categorization:")
        print(f"  Numbered Axioms (001-188): {len(self.numbered_axioms)}")
        print(f"  Prosecution Files: {len(self.prosecution_files)}")
        print(f"  Working Papers: {len(self.working_papers)}")
        print(f"  Case Files: {len(self.case_files)}")
        print(f"  Other Files: {len(self.other_files)}")
        print()
    
    def process_file(self, filepath, file_type):
        """Process a single file and extract all metadata"""
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                content = f.read()
            
            metadata = extract_yaml_frontmatter(content)
            
            # Extract common fields
            data = {
                'File': filepath.name,
                'Path': str(filepath.relative_to(self.axioms_folder)),
                'Type': file_type,
                'Axiom ID': metadata.get('axiom_id', ''),
                'Chain Position': metadata.get('chain_position', ''),
                'Classification': str(metadata.get('classification', '')),
                'Status': metadata.get('status', ''),
                'Stage': metadata.get('stage', ''),
                'Tier': metadata.get('tier', ''),
                'Domain': ', '.join(metadata.get('domain', [])) if isinstance(metadata.get('domain'), list) else metadata.get('domain', ''),
                'Depends On': ', '.join(metadata.get('depends_on', [])) if isinstance(metadata.get('depends_on'), list) else metadata.get('depends_on', ''),
                'Enables': ', '.join(metadata.get('enables', [])) if isinstance(metadata.get('enables'), list) else metadata.get('enables', ''),
                'UUID': metadata.get('uuid', ''),
                'Title': metadata.get('title', ''),
                'Collapse Radius': metadata.get('collapse_radius', ''),
                'Paper Refs': str(metadata.get('paper_refs', [])),
            }
            
            # Extract formal statement
            formal = extract_section(content, 'Formal Statement')
            data['Formal Statement'] = clean_text(formal)
            
            # Extract defeat conditions
            defeat = extract_section(content, 'Defeat Conditions')
            data['Defeat Conditions'] = clean_text(defeat)[:500] if defeat else ''
            
            # Check for prosecution content
            prosecution = extract_section(content, 'Prosecution')
            if prosecution or 'Prosecutorial Analysis' in content:
                data['Has Prosecution'] = 'Yes'
            else:
                data['Has Prosecution'] = 'No'
            
            # Extract objections
            objections = extract_section(content, 'Standard Objections')
            data['Has Objections'] = 'Yes' if objections else 'No'
            
            # Count sections
            section_count = len(re.findall(r'^##\s+', content, re.MULTILINE))
            data['Section Count'] = section_count
            
            # File size
            data['File Size (KB)'] = round(len(content) / 1024, 2)
            
            return data
            
        except Exception as e:
            print(f"Error processing {filepath.name}: {e}")
            return None
    
    def extract_all_data(self):
        """Extract data from all files"""
        print("Extracting data from all files...\n")
        
        all_data = []
        
        # Process numbered axioms
        print(f"Processing {len(self.numbered_axioms)} numbered axioms...")
        for filepath in sorted(self.numbered_axioms):
            data = self.process_file(filepath, 'Numbered Axiom')
            if data:
                all_data.append(data)
        
        # Process prosecution files
        print(f"Processing {len(self.prosecution_files)} prosecution files...")
        for filepath in self.prosecution_files:
            data = self.process_file(filepath, 'Prosecution')
            if data:
                all_data.append(data)
        
        # Process working papers
        print(f"Processing {len(self.working_papers)} working papers...")
        for filepath in self.working_papers:
            data = self.process_file(filepath, 'Working Paper')
            if data:
                all_data.append(data)
        
        # Process case files
        print(f"Processing {len(self.case_files)} case files...")
        for filepath in self.case_files:
            data = self.process_file(filepath, 'Case File')
            if data:
                all_data.append(data)
        
        # Process other files
        print(f"Processing {len(self.other_files)} other files...")
        for filepath in self.other_files:
            data = self.process_file(filepath, 'Other')
            if data:
                all_data.append(data)
        
        print(f"\n✓ Extracted {len(all_data)} files total\n")
        
        return pd.DataFrame(all_data)
    
    def create_comprehensive_excel(self, df, output_file):
        """Create Excel with multiple sheets for different categories"""
        print(f"Creating comprehensive Excel: {output_file}")
        
        wb = Workbook()
        wb.remove(wb.active)
        
        # Define styles
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Sheet 1: All Files
        ws_all = wb.create_sheet("All Files", 0)
        self.write_dataframe_to_sheet(ws_all, df, header_fill, header_font, border)
        
        # Sheet 2: Numbered Axioms Only
        df_numbered = df[df['Type'] == 'Numbered Axiom'].copy()
        if len(df_numbered) > 0:
            df_numbered['Chain Position'] = pd.to_numeric(df_numbered['Chain Position'], errors='coerce')
            df_numbered = df_numbered.sort_values('Chain Position', na_position='last')
            ws_numbered = wb.create_sheet("Numbered Axioms (188)", 1)
            self.write_dataframe_to_sheet(ws_numbered, df_numbered, header_fill, header_font, border)
        
        # Sheet 3: Prosecution Files
        df_prosecution = df[df['Type'] == 'Prosecution'].copy()
        if len(df_prosecution) > 0:
            ws_prosecution = wb.create_sheet("Prosecution", 2)
            self.write_dataframe_to_sheet(ws_prosecution, df_prosecution, header_fill, header_font, border)
        
        # Sheet 4: Working Papers
        df_working = df[df['Type'] == 'Working Paper'].copy()
        if len(df_working) > 0:
            ws_working = wb.create_sheet("Working Papers", 3)
            self.write_dataframe_to_sheet(ws_working, df_working, header_fill, header_font, border)
        
        # Sheet 5: Case Files
        df_case = df[df['Type'] == 'Case File'].copy()
        if len(df_case) > 0:
            ws_case = wb.create_sheet("Case Files", 4)
            self.write_dataframe_to_sheet(ws_case, df_case, header_fill, header_font, border)
        
        # Sheet 6: By Classification
        df_by_class = df[df['Classification'] != ''].copy()
        if len(df_by_class) > 0:
            df_by_class = df_by_class.sort_values('Classification')
            ws_class = wb.create_sheet("By Classification", 5)
            self.write_dataframe_to_sheet(ws_class, df_by_class, header_fill, header_font, border)
        
        # Sheet 7: Summary Statistics
        ws_summary = wb.create_sheet("Summary", 6)
        self.create_summary_sheet(ws_summary, df, header_fill, header_font, border)
        
        wb.save(output_file)
        print(f"✓ Excel file created with {len(wb.sheetnames)} sheets\n")
    
    def write_dataframe_to_sheet(self, ws, df, header_fill, header_font, border):
        """Write DataFrame to worksheet with formatting"""
        # Write headers
        for col_num, column_title in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = column_title
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        
        # Write data
        for row_num, row_data in enumerate(df.values, 2):
            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num)
                if pd.isna(value):
                    cell.value = ""
                elif isinstance(value, (list, dict)):
                    cell.value = str(value)
                else:
                    cell.value = value
                cell.border = border
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                
                # Color code by type
                if col_num == 3:  # Type column
                    if value == 'Numbered Axiom':
                        cell.fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
                    elif value == 'Prosecution':
                        cell.fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
                    elif value == 'Working Paper':
                        cell.fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
                    elif value == 'Case File':
                        cell.fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
        
        # Auto-adjust column widths
        for col_num, column in enumerate(df.columns, 1):
            col_letter = chr(64 + col_num) if col_num <= 26 else chr(64 + col_num // 26) + chr(64 + col_num % 26)
            max_length = max(
                len(str(column)),
                df[column].astype(str).str.len().max() if len(df) > 0 else 0
            )
            ws.column_dimensions[col_letter].width = min(max_length + 2, 60)
        
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    
    def create_summary_sheet(self, ws, df, header_fill, header_font, border):
        """Create summary statistics sheet"""
        summary_data = []
        
        # By type
        type_counts = df['Type'].value_counts()
        for file_type, count in type_counts.items():
            summary_data.append({
                'Category': 'File Type',
                'Value': file_type,
                'Count': count
            })
        
        # By classification
        class_counts = df[df['Classification'] != '']['Classification'].value_counts()
        for classification, count in class_counts.head(20).items():
            summary_data.append({
                'Category': 'Classification',
                'Value': classification,
                'Count': count
            })
        
        # By status
        status_counts = df[df['Status'] != '']['Status'].value_counts()
        for status, count in status_counts.items():
            summary_data.append({
                'Category': 'Status',
                'Value': status,
                'Count': count
            })
        
        # With prosecution
        prosecution_count = len(df[df['Has Prosecution'] == 'Yes'])
        summary_data.append({
            'Category': 'Content',
            'Value': 'Has Prosecution',
            'Count': prosecution_count
        })
        
        # With objections
        objections_count = len(df[df['Has Objections'] == 'Yes'])
        summary_data.append({
            'Category': 'Content',
            'Value': 'Has Objections',
            'Count': objections_count
        })
        
        df_summary = pd.DataFrame(summary_data)
        self.write_dataframe_to_sheet(ws, df_summary, header_fill, header_font, border)

def main():
    axioms_folder = r"O:\Theophysics_Master\TMSUB\GO FOLDER\_AXIOMS_001-188"
    output_file = r"O:\Theophysics_Backend\Python_Backend\Comprehensive_Axiom_Database.xlsx"
    
    print("=" * 80)
    print("COMPREHENSIVE AXIOM SYSTEM EXTRACTOR")
    print("=" * 80)
    print()
    
    extractor = ComprehensiveAxiomExtractor(axioms_folder)
    
    # Scan all files
    extractor.scan_all_files()
    
    # Extract all data
    df = extractor.extract_all_data()
    
    # Create comprehensive Excel
    extractor.create_comprehensive_excel(df, output_file)
    
    # Print statistics
    print("=" * 80)
    print("STATISTICS")
    print("=" * 80)
    print(f"Total Files Processed: {len(df)}")
    print(f"\nBy Type:")
    print(df['Type'].value_counts().to_string())
    print(f"\nBy Classification (top 10):")
    print(df[df['Classification'] != '']['Classification'].value_counts().head(10).to_string())
    print(f"\nFiles with Prosecution: {len(df[df['Has Prosecution'] == 'Yes'])}")
    print(f"Files with Objections: {len(df[df['Has Objections'] == 'Yes'])}")
    
    print("\n" + "=" * 80)
    print("✓ COMPLETE!")
    print("=" * 80)
    print(f"\nOutput: {output_file}")
    print("\nSheets created:")
    print("  1. All Files - Complete database")
    print("  2. Numbered Axioms (188) - Core axioms only")
    print("  3. Prosecution - Prosecution cases")
    print("  4. Working Papers - Working papers")
    print("  5. Case Files - Case files")
    print("  6. By Classification - Sorted by classification")
    print("  7. Summary - Statistics and counts")

if __name__ == "__main__":
    main()
