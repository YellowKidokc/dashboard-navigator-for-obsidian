"""
Prosecution Data to Excel Converter
Extracts prosecution case data and adds it to the Axioms Excel file
"""

import os
import re
from pathlib import Path
import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

def extract_prosecution_data(filepath):
    """Extract prosecution data from YAML file"""
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            data = yaml.safe_load(f)
        
        # Extract strongest support
        support_list = []
        if 'strongest_support' in data and data['strongest_support']:
            for item in data['strongest_support']:
                source = item.get('Source', '')
                argument = item.get('Argument', '')
                support_list.append(f"{source}: {argument}")
        
        # Extract strongest objections
        objection_list = []
        if 'strongest_objections' in data and data['strongest_objections']:
            for item in data['strongest_objections']:
                source = item.get('Source', '')
                objection = item.get('Objection', '')
                objection_list.append(f"{source}: {objection}")
        
        prosecution_data = {
            'File': os.path.basename(filepath),
            'Axiom ID': data.get('axiom_id', ''),
            'Title': data.get('title', ''),
            'Strongest Support': ' | '.join(support_list),
            'Strongest Objections': ' | '.join(objection_list),
            'Verdict': data.get('verdict', ''),
        }
        
        return prosecution_data
    except Exception as e:
        print(f"Error processing {filepath}: {e}")
        return None

def add_prosecution_sheet(excel_file, prosecution_df):
    """Add prosecution data as a new sheet to existing Excel file"""
    wb = load_workbook(excel_file)
    
    # Create new sheet or clear existing
    if "Prosecution" in wb.sheetnames:
        del wb["Prosecution"]
    
    ws = wb.create_sheet("Prosecution", 1)  # Insert as second sheet
    
    # Define styles
    header_fill = PatternFill(start_color="8B0000", end_color="8B0000", fill_type="solid")  # Dark red
    header_font = Font(color="FFFFFF", bold=True, size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Write headers
    for col_num, column_title in enumerate(prosecution_df.columns, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = column_title
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    
    # Write data
    for row_num, row_data in enumerate(prosecution_df.values, 2):
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num)
            if isinstance(value, (list, dict)):
                cell.value = str(value)
            elif pd.isna(value):
                cell.value = ""
            else:
                cell.value = value
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            
            # Highlight verdict column
            if col_num == 6:  # Verdict column
                cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    
    # Adjust column widths
    column_widths = {
        'A': 35,  # File
        'B': 12,  # Axiom ID
        'C': 40,  # Title
        'D': 60,  # Strongest Support
        'E': 60,  # Strongest Objections
        'F': 80,  # Verdict
    }
    
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width
    
    # Freeze header row
    ws.freeze_panes = 'A2'
    
    # Add auto-filter
    ws.auto_filter.ref = ws.dimensions
    
    wb.save(excel_file)
    print(f"✓ Added 'Prosecution' sheet to Excel file")

def main():
    """Main execution function"""
    prosecuted_folder = r"O:\Theophysics_Master\TMSUB\GO FOLDER\_AXIOMS_001-188\_PROSECUTED"
    excel_file = r"O:\Theophysics_Backend\Python_Backend\Axioms_Database.xlsx"
    
    print("=" * 70)
    print("PROSECUTION DATA TO EXCEL CONVERTER")
    print("=" * 70)
    print(f"\nSource folder: {prosecuted_folder}")
    print(f"Excel file: {excel_file}\n")
    
    # Get all YAML files
    yaml_files = sorted([f for f in Path(prosecuted_folder).glob("*.yaml") if f.is_file()])
    print(f"Found {len(yaml_files)} prosecution files\n")
    
    # Process each file
    prosecution_data = []
    for idx, filepath in enumerate(yaml_files, 1):
        print(f"Processing [{idx}/{len(yaml_files)}]: {filepath.name}")
        data = extract_prosecution_data(filepath)
        if data:
            prosecution_data.append(data)
    
    # Create DataFrame
    df = pd.DataFrame(prosecution_data)
    
    print(f"\n✓ Processed {len(prosecution_data)} prosecution cases successfully")
    print(f"\nAdding to Excel file...")
    
    # Add to Excel
    add_prosecution_sheet(excel_file, df)
    
    # Print summary
    print("\n" + "=" * 70)
    print("SUMMARY")
    print("=" * 70)
    print(f"Total Prosecution Cases: {len(df)}")
    print(f"Axioms Prosecuted: {', '.join(df['Axiom ID'].tolist())}")
    print("\n" + "=" * 70)
    print("✓ COMPLETE!")
    print("=" * 70)

if __name__ == "__main__":
    main()
