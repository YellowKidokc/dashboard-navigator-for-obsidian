"""
Unfulfilled Prophecy Aggregator
Filters unfulfilled prophecies and aggregates research links, theories, and cross-references
Generates comprehensive research database for each unfulfilled prophecy
"""

import pandas as pd
import os
import re
from pathlib import Path
from collections import defaultdict
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

class UnfulfilledProphecyAggregator:
    def __init__(self, excel_file):
        self.excel_file = excel_file
        self.data = None
        self.unfulfilled = []
        self.research_links = {}
        
    def load_data(self):
        """Load Excel file"""
        print(f"Loading: {self.excel_file}\n")
        
        try:
            df = pd.read_excel(self.excel_file, sheet_name=0)
            self.data = df
            print(f"✓ Loaded {len(df)} prophecies\n")
            return True
        except Exception as e:
            print(f"Error loading file: {e}")
            return False
    
    def filter_unfulfilled(self):
        """Filter for unfulfilled prophecies"""
        print("Filtering unfulfilled prophecies...")
        
        # Check column names
        status_col = None
        for col in self.data.columns:
            if 'status' in col.lower() or 'fulfilled' in col.lower() or 'column4' in col.lower():
                status_col = col
                break
        
        if status_col is None:
            # Try to find by position (4th column is status)
            status_col = self.data.columns[3] if len(self.data.columns) > 3 else None
        
        if status_col:
            # First, show what status values exist
            print(f"Status column: {status_col}")
            status_values = self.data[status_col].value_counts()
            print(f"Status values found:")
            for val, count in status_values.head(10).items():
                print(f"  '{val}': {count}")
            print()
            
            # Filter for NOT fulfilled (more flexible)
            # Look for anything that's NOT explicitly "Fulfilled"
            mask = ~self.data[status_col].astype(str).str.lower().str.contains('fulfilled', na=False)
            # Also exclude NaN/empty
            mask = mask & self.data[status_col].notna()
            
            # Alternative: look for specific unfulfilled markers
            unfulfilled_markers = ['unfulfilled', 'pending', 'future', 'not yet', 'awaiting', 'partial']
            unfulfilled_mask = self.data[status_col].astype(str).str.lower().str.contains('|'.join(unfulfilled_markers), na=False)
            
            # Combine: either explicitly unfulfilled OR not explicitly fulfilled
            final_mask = unfulfilled_mask | (mask & ~self.data[status_col].astype(str).str.lower().str.contains('fulfilled', na=False))
            
            self.unfulfilled = self.data[final_mask].copy()
            
            print(f"✓ Found {len(self.unfulfilled)} unfulfilled/pending prophecies\n")
        else:
            print("⚠️  Could not identify status column")
            self.unfulfilled = pd.DataFrame()
        
        return len(self.unfulfilled)
    
    def generate_research_links(self, prophecy_text, book, reference):
        """Generate research links for a prophecy"""
        links = {}
        
        # Clean text for search
        search_text = re.sub(r'[^\w\s]', ' ', prophecy_text)
        search_terms = ' '.join(search_text.split()[:10])  # First 10 words
        
        # Bible Gateway
        bible_ref = f"{book} {reference}".replace(' ', '+')
        links['Bible Gateway'] = f"https://www.biblegateway.com/passage/?search={bible_ref}&version=ESV"
        
        # Blue Letter Bible
        links['Blue Letter Bible'] = f"https://www.blueletterbible.org/search/search.cfm?Criteria={bible_ref}"
        
        # Google Scholar
        scholar_query = f"{book}+{reference}+prophecy+unfulfilled".replace(' ', '+')
        links['Google Scholar'] = f"https://scholar.google.com/scholar?q={scholar_query}"
        
        # Bible Hub
        links['Bible Hub'] = f"https://biblehub.com/commentaries/{book.lower()}/{reference.split(':')[0]}.htm"
        
        # Got Questions
        got_q = search_terms.replace(' ', '+')
        links['Got Questions'] = f"https://www.gotquestions.org/search?query={got_q}"
        
        # Bible.org
        links['Bible.org'] = f"https://bible.org/search/apachesolr_search/{search_terms.replace(' ', '%20')}"
        
        # Theology search
        links['Theological Studies'] = f"https://www.google.com/search?q={search_terms.replace(' ', '+')}+theology+eschatology"
        
        return links
    
    def generate_theories_context(self, prophecy_text, notes):
        """Generate theories and interpretations context"""
        theories = []
        
        # Extract from notes if available
        if pd.notna(notes):
            notes_str = str(notes)
            
            # Look for citations
            citations = re.findall(r'\[cite:.*?\]', notes_str)
            if citations:
                theories.append(f"Citations found: {', '.join(citations)}")
            
            # Look for interpretations
            if 'interpret' in notes_str.lower():
                theories.append("Multiple interpretations noted")
            
            # Look for eschatological markers
            eschatology_terms = ['end times', 'second coming', 'millennium', 'tribulation', 'rapture']
            found_terms = [term for term in eschatology_terms if term in notes_str.lower()]
            if found_terms:
                theories.append(f"Eschatological: {', '.join(found_terms)}")
        
        # Analyze prophecy text
        text_lower = prophecy_text.lower()
        
        # Messianic indicators
        messianic_terms = ['messiah', 'christ', 'son of', 'king', 'lord', 'savior']
        if any(term in text_lower for term in messianic_terms):
            theories.append("Potential Messianic prophecy")
        
        # End times indicators
        end_times_terms = ['last days', 'end of', 'final', 'judgment', 'return', 'coming']
        if any(term in text_lower for term in end_times_terms):
            theories.append("End times prophecy")
        
        # Israel/Nations
        if 'israel' in text_lower or 'jerusalem' in text_lower or 'nations' in text_lower:
            theories.append("Israel/Nations prophecy")
        
        if not theories:
            theories.append("Requires theological analysis")
        
        return theories
    
    def aggregate_all_unfulfilled(self):
        """Aggregate research for all unfulfilled prophecies"""
        print("Aggregating research links and theories...\n")
        
        aggregated = []
        
        for idx, row in self.unfulfilled.iterrows():
            # Extract data
            book = str(row.iloc[0]) if len(row) > 0 else "Unknown"
            reference = str(row.iloc[1]) if len(row) > 1 else ""
            prophecy = str(row.iloc[2]) if len(row) > 2 else ""
            status = str(row.iloc[3]) if len(row) > 3 else ""
            notes = str(row.iloc[4]) if len(row) > 4 else ""
            
            # Generate research links
            links = self.generate_research_links(prophecy, book, reference)
            
            # Generate theories
            theories = self.generate_theories_context(prophecy, notes)
            
            # Compile
            aggregated.append({
                'Book': book,
                'Reference': reference,
                'Prophecy': prophecy,
                'Status': status,
                'Notes': notes,
                'Theories': ' | '.join(theories),
                'Bible Gateway': links['Bible Gateway'],
                'Blue Letter Bible': links['Blue Letter Bible'],
                'Bible Hub': links['Bible Hub'],
                'Google Scholar': links['Google Scholar'],
                'Got Questions': links['Got Questions'],
                'Bible.org': links['Bible.org'],
                'Theological Studies': links['Theological Studies']
            })
        
        print(f"✓ Aggregated research for {len(aggregated)} unfulfilled prophecies\n")
        
        return pd.DataFrame(aggregated)
    
    def export_to_excel(self, df, output_file):
        """Export aggregated data to Excel with hyperlinks"""
        print(f"Exporting to: {output_file}")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Unfulfilled Prophecies"
        
        # Define styles
        header_fill = PatternFill(start_color="8B0000", end_color="8B0000", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        link_font = Font(color="0000FF", underline="single")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Write headers
        for col_num, column_title in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = column_title
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        
        # Write data
        for row_num, row_data in enumerate(df.values, 2):
            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num)
                
                # Check if this is a link column
                col_name = df.columns[col_num - 1]
                is_link = any(x in col_name for x in ['Gateway', 'Bible', 'Scholar', 'Questions', 'Theological', 'Hub'])
                
                if is_link and pd.notna(value) and str(value).startswith('http'):
                    # Create hyperlink
                    cell.hyperlink = value
                    cell.value = col_name  # Show column name as link text
                    cell.font = link_font
                else:
                    cell.value = value if pd.notna(value) else ""
                
                cell.border = border
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        
        # Adjust column widths
        column_widths = {
            'A': 15,  # Book
            'B': 12,  # Reference
            'C': 50,  # Prophecy
            'D': 12,  # Status
            'E': 40,  # Notes
            'F': 40,  # Theories
            'G': 15,  # Bible Gateway
            'H': 18,  # Blue Letter Bible
            'I': 15,  # Bible Hub
            'J': 15,  # Google Scholar
            'K': 15,  # Got Questions
            'L': 12,  # Bible.org
            'M': 18,  # Theological Studies
        }
        
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width
        
        # Freeze header and first 3 columns
        ws.freeze_panes = 'D2'
        
        # Add auto-filter
        ws.auto_filter.ref = ws.dimensions
        
        wb.save(output_file)
        print(f"✓ Excel file saved with {len(df)} unfulfilled prophecies\n")
    
    def generate_summary_report(self, df):
        """Generate text summary report"""
        report = []
        report.append("=" * 80)
        report.append("UNFULFILLED PROPHECY RESEARCH AGGREGATION REPORT")
        report.append("=" * 80)
        report.append(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        report.append("")
        
        report.append("SUMMARY:")
        report.append(f"  Total Unfulfilled Prophecies: {len(df)}")
        report.append("")
        
        # Group by book
        by_book = df.groupby('Book').size().sort_values(ascending=False)
        report.append("UNFULFILLED PROPHECIES BY BOOK:")
        for book, count in by_book.items():
            report.append(f"  {book}: {count}")
        report.append("")
        
        # Theory categories
        theory_counts = defaultdict(int)
        for theories in df['Theories']:
            for theory in str(theories).split(' | '):
                theory_counts[theory] += 1
        
        report.append("THEORY CATEGORIES:")
        for theory, count in sorted(theory_counts.items(), key=lambda x: -x[1]):
            report.append(f"  {theory}: {count}")
        report.append("")
        
        # Sample entries
        report.append("SAMPLE UNFULFILLED PROPHECIES:")
        for idx, row in df.head(10).iterrows():
            report.append(f"\n{row['Book']} {row['Reference']}:")
            report.append(f"  Prophecy: {row['Prophecy'][:100]}...")
            report.append(f"  Theories: {row['Theories']}")
        
        report.append("")
        report.append("=" * 80)
        
        return "\n".join(report)

def main():
    excel_file = r"C:\Users\lowes\Downloads\Bible prophecies full.xlsx"
    output_dir = r"O:\Theophysics_Backend\Python_Backend"
    
    print("=" * 80)
    print("UNFULFILLED PROPHECY RESEARCH AGGREGATOR")
    print("=" * 80)
    print()
    
    aggregator = UnfulfilledProphecyAggregator(excel_file)
    
    # Load data
    if not aggregator.load_data():
        return
    
    # Filter unfulfilled
    count = aggregator.filter_unfulfilled()
    if count == 0:
        print("No unfulfilled prophecies found.")
        return
    
    # Aggregate research
    df_aggregated = aggregator.aggregate_all_unfulfilled()
    
    # Export to Excel
    output_excel = os.path.join(output_dir, "unfulfilled_prophecies_research.xlsx")
    aggregator.export_to_excel(df_aggregated, output_excel)
    
    # Generate report
    report = aggregator.generate_summary_report(df_aggregated)
    print(report)
    
    # Save report
    report_file = os.path.join(output_dir, "unfulfilled_prophecies_report.txt")
    with open(report_file, 'w', encoding='utf-8') as f:
        f.write(report)
    
    print(f"\n✓ Report saved: {report_file}")
    
    print("\n" + "=" * 80)
    print("✓ AGGREGATION COMPLETE!")
    print("=" * 80)
    print(f"\nOutputs:")
    print(f"  - Excel with Research Links: {output_excel}")
    print(f"  - Summary Report: {report_file}")
    print(f"\nThe Excel file contains clickable hyperlinks to:")
    print(f"  - Bible Gateway (ESV)")
    print(f"  - Blue Letter Bible")
    print(f"  - Bible Hub Commentaries")
    print(f"  - Google Scholar (academic research)")
    print(f"  - Got Questions (apologetics)")
    print(f"  - Bible.org (theological studies)")
    print(f"  - General theological research")

if __name__ == "__main__":
    main()
