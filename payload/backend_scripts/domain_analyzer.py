"""
Domain Analyzer
Groups and analyzes axioms by domain (physics, theology, consciousness, etc.)
Generates domain-specific reports and cross-domain bridge analysis
"""

import os
import re
from pathlib import Path
import yaml
import pandas as pd
from collections import defaultdict

def extract_yaml_frontmatter(content):
    """Extract YAML frontmatter from markdown content"""
    pattern = r'^---\s*\n(.*?)\n---\s*\n'
    match = re.match(pattern, content, re.DOTALL)
    if match:
        try:
            return yaml.safe_load(match.group(1))
        except yaml.YAMLError:
            return {}
    return {}

class DomainAnalyzer:
    def __init__(self, axioms_folder):
        self.axioms_folder = axioms_folder
        self.axioms = {}
        self.domain_groups = defaultdict(list)
        self.cross_domain = []
        
    def load_axioms(self):
        """Load all axiom files"""
        md_files = []
        for f in Path(self.axioms_folder).glob("*.md"):
            if f.is_file() and re.match(r'^\d{3}_', f.name):
                md_files.append(f)
        
        print(f"Loading {len(md_files)} axiom files...")
        
        for filepath in md_files:
            try:
                with open(filepath, 'r', encoding='utf-8') as f:
                    content = f.read()
                
                metadata = extract_yaml_frontmatter(content)
                axiom_id = metadata.get('axiom_id', '')
                
                if axiom_id:
                    domains = metadata.get('domain', [])
                    if isinstance(domains, str):
                        domains = [domains]
                    elif not domains:
                        domains = ['uncategorized']
                    
                    self.axioms[axiom_id] = {
                        'file': filepath.name,
                        'chain_position': metadata.get('chain_position', 999),
                        'domains': domains,
                        'classification': str(metadata.get('classification', '')),
                        'status': metadata.get('status', ''),
                        'stage': metadata.get('stage', 0),
                        'tier': metadata.get('tier', 0)
                    }
                    
                    # Group by domain
                    for domain in domains:
                        self.domain_groups[domain].append(axiom_id)
                    
                    # Track cross-domain axioms
                    if len(domains) > 1:
                        self.cross_domain.append({
                            'axiom_id': axiom_id,
                            'domains': domains,
                            'chain_position': metadata.get('chain_position', 999)
                        })
                        
            except Exception as e:
                print(f"Error loading {filepath.name}: {e}")
        
        print(f"✓ Loaded {len(self.axioms)} axioms\n")
    
    def generate_domain_report(self):
        """Generate comprehensive domain analysis report"""
        report = []
        report.append("=" * 80)
        report.append("DOMAIN ANALYSIS REPORT")
        report.append("=" * 80)
        report.append("")
        
        # Overall statistics
        report.append("OVERALL STATISTICS:")
        report.append(f"  Total Axioms: {len(self.axioms)}")
        report.append(f"  Total Domains: {len(self.domain_groups)}")
        report.append(f"  Cross-Domain Axioms: {len(self.cross_domain)}")
        report.append("")
        
        # Domain breakdown
        report.append("=" * 80)
        report.append("AXIOMS BY DOMAIN:")
        report.append("=" * 80)
        
        for domain, axioms in sorted(self.domain_groups.items(), key=lambda x: -len(x[1])):
            report.append(f"\n{domain.upper()} ({len(axioms)} axioms):")
            report.append("-" * 40)
            
            # Group by stage within domain
            by_stage = defaultdict(list)
            for axiom_id in axioms:
                stage = self.axioms[axiom_id]['stage']
                by_stage[stage].append(axiom_id)
            
            for stage in sorted(by_stage.keys()):
                stage_axioms = sorted(by_stage[stage], key=lambda x: self.axioms[x]['chain_position'])
                report.append(f"  Stage {stage}: {', '.join(stage_axioms)}")
        
        report.append("")
        
        # Cross-domain bridges
        report.append("=" * 80)
        report.append("CROSS-DOMAIN BRIDGES:")
        report.append("=" * 80)
        report.append(f"\n{len(self.cross_domain)} axioms span multiple domains:\n")
        
        for item in sorted(self.cross_domain, key=lambda x: x['chain_position']):
            domains_str = ' + '.join(item['domains'])
            report.append(f"  {item['axiom_id']} (pos {item['chain_position']}): {domains_str}")
        
        report.append("")
        
        # Domain interaction matrix
        report.append("=" * 80)
        report.append("DOMAIN INTERACTION MATRIX:")
        report.append("=" * 80)
        report.append("")
        
        domain_pairs = defaultdict(int)
        for item in self.cross_domain:
            domains = sorted(item['domains'])
            for i in range(len(domains)):
                for j in range(i+1, len(domains)):
                    pair = f"{domains[i]} ↔ {domains[j]}"
                    domain_pairs[pair] += 1
        
        for pair, count in sorted(domain_pairs.items(), key=lambda x: -x[1]):
            report.append(f"  {pair}: {count} axioms")
        
        report.append("")
        report.append("=" * 80)
        
        return "\n".join(report)
    
    def generate_domain_excel(self, output_file):
        """Generate Excel file with domain-specific sheets"""
        print("Generating domain-specific Excel file...")
        
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet
        
        # Define styles
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Create a sheet for each domain
        for domain, axiom_ids in sorted(self.domain_groups.items(), key=lambda x: -len(x[1])):
            # Prepare data
            domain_data = []
            for axiom_id in sorted(axiom_ids, key=lambda x: self.axioms[x]['chain_position']):
                data = self.axioms[axiom_id]
                domain_data.append({
                    'Axiom ID': axiom_id,
                    'Chain Position': data['chain_position'],
                    'Status': data['status'],
                    'Stage': data['stage'],
                    'Tier': data['tier'],
                    'All Domains': ', '.join(data['domains']),
                    'File': data['file']
                })
            
            df = pd.DataFrame(domain_data)
            
            # Create sheet (Excel sheet names max 31 chars)
            sheet_name = domain[:31]
            ws = wb.create_sheet(sheet_name)
            
            # Write headers
            for col_num, column_title in enumerate(df.columns, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.value = column_title
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = border
            
            # Write data
            for row_num, row_data in enumerate(df.values, 2):
                for col_num, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_num, column=col_num)
                    if pd.isna(value):
                        cell.value = ""
                    else:
                        cell.value = value
                    cell.border = border
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            
            # Adjust column widths
            ws.column_dimensions['A'].width = 15
            ws.column_dimensions['B'].width = 15
            ws.column_dimensions['C'].width = 20
            ws.column_dimensions['D'].width = 10
            ws.column_dimensions['E'].width = 10
            ws.column_dimensions['F'].width = 30
            ws.column_dimensions['G'].width = 35
            
            ws.freeze_panes = 'A2'
        
        # Add summary sheet
        ws_summary = wb.create_sheet("Summary", 0)
        summary_data = []
        for domain, axioms in sorted(self.domain_groups.items(), key=lambda x: -len(x[1])):
            summary_data.append({
                'Domain': domain,
                'Count': len(axioms),
                'Percentage': f"{len(axioms)/len(self.axioms)*100:.1f}%"
            })
        
        df_summary = pd.DataFrame(summary_data)
        
        for col_num, column_title in enumerate(df_summary.columns, 1):
            cell = ws_summary.cell(row=1, column=col_num)
            cell.value = column_title
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        for row_num, row_data in enumerate(df_summary.values, 2):
            for col_num, value in enumerate(row_data, 1):
                cell = ws_summary.cell(row=row_num, column=col_num)
                cell.value = value
                cell.border = border
                cell.alignment = Alignment(vertical='top')
        
        ws_summary.column_dimensions['A'].width = 25
        ws_summary.column_dimensions['B'].width = 15
        ws_summary.column_dimensions['C'].width = 15
        
        wb.save(output_file)
        print(f"✓ Domain Excel file saved: {output_file}\n")

def main():
    axioms_folder = r"O:\Theophysics_Master\TMSUB\GO FOLDER\_AXIOMS_001-188"
    output_dir = r"O:\Theophysics_Backend\Python_Backend"
    
    print("=" * 80)
    print("DOMAIN ANALYZER")
    print("=" * 80)
    print(f"\nSource: {axioms_folder}\n")
    
    analyzer = DomainAnalyzer(axioms_folder)
    analyzer.load_axioms()
    
    # Generate report
    report = analyzer.generate_domain_report()
    print(report)
    
    # Save report
    report_file = os.path.join(output_dir, "domain_analysis_report.txt")
    with open(report_file, 'w', encoding='utf-8') as f:
        f.write(report)
    print(f"\n✓ Report saved: {report_file}")
    
    # Generate Excel
    excel_file = os.path.join(output_dir, "axioms_by_domain.xlsx")
    analyzer.generate_domain_excel(excel_file)
    
    print(f"\n✓ Domain analysis complete!")

if __name__ == "__main__":
    main()
