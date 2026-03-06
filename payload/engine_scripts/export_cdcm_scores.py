#!/usr/bin/env python3
"""
CDCM SCORE EXPORTER
Export JSON scoring results to Excel matching CDCM_final rubric format.

Usage:
  python export_cdcm_scores.py score_output.json -o output.xlsx
"""

import sys
import json
from pathlib import Path
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

# Criterion metadata
CRITERIA = {
    "A": {"name": "CITATION QUALITY", "color": "FFE6E6"},
    "B": {"name": "FACTUAL ACCURACY", "color": "FFE6F0"},
    "C": {"name": "LOGICAL RIGOR", "color": "FFE6FF"},
    "D": {"name": "EVIDENCE SUFFICIENCY", "color": "F0E6FF"},
    "E": {"name": "SEMANTIC PRECISION", "color": "E6F0FF"},
    "F": {"name": "ACADEMIC CONVENTION", "color": "E6FFFF"},
    "G": {"name": "WORDING & CLARITY", "color": "E6FFF0"},
    "H": {"name": "CROSS-DOMAIN VALIDITY", "color": "F0FFE6"},
    "I": {"name": "FALSIFIABILITY & SCOPE", "color": "FFF0E6"},
    "J": {"name": "EXTERNAL THEORY USAGE", "color": "FFFFE6"},
    "K": {"name": "NOVELTY ASSESSMENT", "color": "F0F0E6"},
}

CRITERIA_DETAILS = {
    "A1": "Source Existence & Verifiability",
    "A2": "Source Authority & Relevance",
    "A3": "Citation Formatting",
    "A4": "Citation Density & Distribution",
    "B1": "Empirical Claim Accuracy",
    "B2": "Historical/Contextual Accuracy",
    "B3": "Mathematical/Formal Accuracy",
    "B4": "Claim-Evidence Alignment",
    "C1": "Deductive Validity",
    "C2": "Inductive Strength",
    "C3": "Absence of Fallacies",
    "C4": "Argument Chain Integrity",
    "D1": "Claim Strength vs Evidence Weight",
    "D2": "Counter-Evidence Engagement",
    "D3": "Replication & Corroboration",
    "D4": "Statistical Rigor",
    "E1": "Term Definition Consistency",
    "E2": "Equivocation Detection",
    "E3": "Domain Translation Fidelity",
    "E4": "Hedging Calibration",
    "F1": "Abstract Quality",
    "F2": "Methodology Transparency",
    "F3": "Structure & Flow",
    "F4": "Peer Convention Compliance",
    "G1": "Prose Precision",
    "G2": "Jargon Accessibility",
    "G3": "Concision",
    "G4": "Tone Calibration",
    "H1": "Isomorphism vs Analogy Distinction",
    "H2": "Domain Boundary Respect",
    "H3": "Mapping Rigor",
    "H4": "Prediction Transfer",
    "I1": "Falsification Criteria Stated",
    "I2": "Scope Boundaries Declared",
    "I3": "Failure Mode Acknowledgment",
    "I4": "Testable Prediction Specificity",
    "J1": "Theory Representation Accuracy",
    "J2": "Scope Boundary Respect for Borrowed Theories",
    "J3": "Integration vs Appropriation",
    "J4": "Competing Interpretation Acknowledgment",
    "K1": "Conceptual Originality",
    "K2": "Nearest Existing Framework Distance",
    "K3": "Methodological Innovation",
    "K4": "Gap-Filling vs Gap-Creating",
}

FIX_TYPES = {
    "A": "CITE",
    "B": "FACT-CHECK",
    "C": "LOGIC",
    "D": "EVIDENCE",
    "E": "REWORD",
    "F": "RESTRUCTURE",
    "G": "REWRITE",
    "H": "REMAP",
    "I": "SPECIFY",
    "J": "VERIFY-THEORY",
    "K": "NOVELTY",
}

FIX_PRIORITY = {
    "B": "P1-CRITICAL",
    "C": "P1-CRITICAL",
    "J": "P1-CRITICAL",
    "A": "P2-HIGH",
    "D": "P2-HIGH",
    "E": "P2-HIGH",
    "H": "P2-HIGH",
    "K": "P2-HIGH",
    "F": "P3-MEDIUM",
    "G": "P3-MEDIUM",
    "I": "P3-MEDIUM",
}

def export_to_excel(json_file: Path, output_file: Path):
    """Convert JSON scores to Excel workbook."""

    with open(json_file) as f:
        data = json.load(f)

    scores = data.get("scores", {})
    composite = data.get("composite_score", 0)
    metadata = data.get("metadata", {})
    flagged = data.get("flagged_criteria", [])

    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CDCM Scores"

    # Header
    ws["A1"] = "CDCM PAPER SCORING REPORT"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:H1")

    # Metadata row
    row = 3
    ws[f"A{row}"] = f"Paper: {metadata.get('paper', '?')}"
    ws[f"A{row+1}"] = f"Date: {metadata.get('timestamp', '?')[:10]}"
    ws[f"A{row+2}"] = f"Scorer: {metadata.get('scorer', '?')}"

    # Composite score
    row = 7
    ws[f"A{row}"] = "COMPOSITE SCORE"
    ws[f"B{row}"] = composite
    ws[f"B{row}"].font = Font(bold=True, size=12)

    # Column headers
    row = 9
    headers = ["#", "CRITERION", "WHAT IT CATCHES", "SCORE (0-10)", "FLAG", "FIX TYPE", "PRIORITY", "NOTES"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row, col)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")

    # Scores rows
    row = 10
    for crit_id in sorted(scores.keys()):
        dim = crit_id[0]
        score = scores[crit_id]

        # Find flag
        flag_info = next((f for f in flagged if f["criterion"] == crit_id), None)
        is_flagged = score < 7

        ws.cell(row, 1).value = crit_id
        ws.cell(row, 2).value = CRITERIA_DETAILS.get(crit_id, "")
        ws.cell(row, 3).value = ""  # WHAT IT CATCHES - from original rubric
        ws.cell(row, 4).value = score
        ws.cell(row, 5).value = "🚩" if is_flagged else ""
        ws.cell(row, 6).value = flag_info["fix_type"] if flag_info else ""
        ws.cell(row, 7).value = flag_info["priority"] if flag_info else ""
        ws.cell(row, 8).value = ""  # NOTES placeholder

        # Color code by dimension
        color = CRITERIA[dim]["color"]
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        for col in range(1, 9):
            ws.cell(row, col).fill = fill

        # Flag color
        if is_flagged:
            ws.cell(row, 5).fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            ws.cell(row, 5).font = Font(color="FFFFFF", bold=True)

        row += 1

    # Summary section
    row += 2
    ws[f"A{row}"] = "FLAGGED CRITERIA (<7)"
    ws[f"A{row}"].font = Font(bold=True)

    row += 1
    for flag_info in flagged:
        ws[f"A{row}"] = flag_info["criterion"]
        ws[f"B{row}"] = flag_info["name"]
        ws[f"C{row}"] = flag_info["score"]
        ws[f"D{row}"] = flag_info["fix_type"]
        ws[f"E{row}"] = flag_info["priority"]
        row += 1

    # Adjust column widths
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 15
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 30

    wb.save(output_file)
    print(f"[OK] Exported to: {output_file}")

def main():
    if len(sys.argv) < 2:
        print("Usage: python export_cdcm_scores.py score_output.json -o output.xlsx")
        sys.exit(1)

    json_file = Path(sys.argv[1])
    output_file = None

    if "-o" in sys.argv:
        output_file = Path(sys.argv[sys.argv.index("-o") + 1])
    else:
        output_file = json_file.parent / f"{json_file.stem}.xlsx"

    if not json_file.exists():
        print(f"[ERROR] File not found: {json_file}")
        sys.exit(1)

    export_to_excel(json_file, output_file)

if __name__ == "__main__":
    main()
