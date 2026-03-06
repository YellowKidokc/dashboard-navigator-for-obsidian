"""
Run this script manually in a terminal:
  cd O:\_Theophysics_v3
  python _run_me.py

It will:
1. Read the Excel file
2. Write the data to _excel_output.json
"""
import json
import sys
import os

try:
    import openpyxl
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    import openpyxl

excel_path = r"C:\Users\lowes\Downloads\landscape_consciousness_theories_links.xlsx"
output_path = r"O:\_Theophysics_v3\_excel_output.json"

print(f"Reading: {excel_path}")
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

headers = [cell.value for cell in ws[1]]
print(f"Headers: {headers}")
print(f"Max row: {ws.max_row}, Max col: {ws.max_column}")

rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    row_data = {}
    for i, val in enumerate(row):
        if i < len(headers) and val is not None:
            row_data[str(headers[i])] = str(val)
    if row_data:
        rows.append(row_data)

print(f"Total data rows: {len(rows)}")
for i, r in enumerate(rows):
    print(f"  Row {i+1}: {r}")

result = {"headers": headers, "total_rows": len(rows), "rows": rows}
with open(output_path, "w", encoding="utf-8") as f:
    json.dump(result, f, indent=2, ensure_ascii=False)

print(f"\nJSON written to: {output_path}")
print("DONE")
