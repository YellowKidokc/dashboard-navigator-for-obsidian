"""
Create comprehensive Excel report for Adam and Eve papers scoring
"""

import pandas as pd
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

def create_excel_report():
    """Create formatted Excel report"""
    
    outputs_dir = Path(r"O:\_Theophysics_v3\00_SYSTEM\01_ENGINE\scripts\from_Note\outputs")
    
    # Load comparison CSV
    comparison_df = pd.read_csv(outputs_dir / "adam_eve_comparison.csv")
    
    # Load all JSON scores
    json_files = list(outputs_dir.glob("*_scores.json"))
    all_scores = {}
    
    for json_file in json_files:
        with open(json_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
            title = json_file.stem.replace('_scores', '')
            all_scores[title] = data
    
    # Create workbook
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet
    
    # ========================================================================
    # SHEET 1: Summary Comparison
    # ========================================================================
    ws_summary = wb.create_sheet("Summary")
    
    # Title
    ws_summary['A1'] = "Adam and Eve Papers - Unified Scoring Analysis"
    ws_summary['A1'].font = Font(size=16, bold=True, color="FFFFFF")
    ws_summary['A1'].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    ws_summary.merge_cells('A1:L1')
    
    # Headers
    headers = ['Rank', 'Paper Title', 'χ (Chi)', 'κ (Kappa)', 'ρ (Rho)', 
               'Π (Pi)', 'A (Anthropos)', 'Λ (Lambda)', 
               'Metrics', 'Evidence', 'Vetoes', 'Warnings']
    
    for col, header in enumerate(headers, 1):
        cell = ws_summary.cell(row=3, column=col)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    # Data rows
    for idx, row in comparison_df.iterrows():
        excel_row = idx + 4
        ws_summary.cell(row=excel_row, column=1, value=row['Rank'])
        ws_summary.cell(row=excel_row, column=2, value=row['Title'])
        ws_summary.cell(row=excel_row, column=3, value=row['Chi'])
        ws_summary.cell(row=excel_row, column=4, value=row['Kappa'])
        ws_summary.cell(row=excel_row, column=5, value=row['Rho'])
        ws_summary.cell(row=excel_row, column=6, value=row['Pi'])
        ws_summary.cell(row=excel_row, column=7, value=row['A'])
        ws_summary.cell(row=excel_row, column=8, value=row['Lambda'])
        ws_summary.cell(row=excel_row, column=9, value=row['Metrics'])
        ws_summary.cell(row=excel_row, column=10, value=row['Evidence'])
        ws_summary.cell(row=excel_row, column=11, value=row['Vetoes'])
        ws_summary.cell(row=excel_row, column=12, value=row['Warnings'])
        
        # Color code Chi scores
        chi_cell = ws_summary.cell(row=excel_row, column=3)
        if row['Chi'] >= 7:
            chi_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        elif row['Chi'] >= 5:
            chi_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        else:
            chi_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        
        # Highlight vetoes
        if row['Vetoes'] > 0:
            ws_summary.cell(row=excel_row, column=11).font = Font(color="FF0000", bold=True)
    
    # Statistics row
    stats_row = len(comparison_df) + 5
    ws_summary.cell(row=stats_row, column=1, value="AVERAGE")
    ws_summary.cell(row=stats_row, column=1).font = Font(bold=True)
    ws_summary.cell(row=stats_row, column=3, value=comparison_df['Chi'].mean())
    ws_summary.cell(row=stats_row, column=4, value=comparison_df['Kappa'].mean())
    ws_summary.cell(row=stats_row, column=5, value=comparison_df['Rho'].mean())
    
    # Column widths
    ws_summary.column_dimensions['A'].width = 8
    ws_summary.column_dimensions['B'].width = 30
    for col in ['C', 'D', 'E', 'F', 'G', 'H']:
        ws_summary.column_dimensions[col].width = 12
    for col in ['I', 'J', 'K', 'L']:
        ws_summary.column_dimensions[col].width = 10
    
    # ========================================================================
    # SHEET 2: Fruit Scores Matrix
    # ========================================================================
    ws_fruits = wb.create_sheet("Fruits Matrix")
    
    # Title
    ws_fruits['A1'] = "12 Fruits of the Spirit - Scores Across Papers"
    ws_fruits['A1'].font = Font(size=14, bold=True, color="FFFFFF")
    ws_fruits['A1'].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    ws_fruits.merge_cells('A1:H1')
    
    # Build fruit matrix
    fruit_names = [
        'Love', 'Joy', 'Peace', 'Patience', 'Kindness', 'Goodness',
        'Faithfulness', 'Gentleness', 'Self-Control', 'Truth', 'Wisdom', 'Grace'
    ]
    
    # Headers
    ws_fruits['A3'] = "Fruit"
    ws_fruits['A3'].font = Font(bold=True)
    
    for col_idx, title in enumerate(sorted(all_scores.keys()), 2):
        cell = ws_fruits.cell(row=3, column=col_idx)
        cell.value = title[:20]  # Truncate long titles
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", text_rotation=45)
    
    avg_col = len(all_scores) + 2
    ws_fruits.cell(row=3, column=avg_col, value="Average")
    ws_fruits.cell(row=3, column=avg_col).font = Font(bold=True)
    
    # Fruit rows
    for fruit_idx, fruit_name in enumerate(fruit_names, 4):
        ws_fruits.cell(row=fruit_idx, column=1, value=fruit_name)
        
        fruit_scores = []
        for col_idx, (title, data) in enumerate(sorted(all_scores.items()), 2):
            # Find matching fruit
            net_score = 0
            for fruit in data['fruits']:
                if fruit['name'] == fruit_name:
                    net_score = fruit['net']
                    break
            
            cell = ws_fruits.cell(row=fruit_idx, column=col_idx)
            cell.value = net_score
            fruit_scores.append(net_score)
            
            # Color code
            if net_score > 0.1:
                cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            elif net_score < -0.1:
                cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        
        # Average
        avg = sum(fruit_scores) / len(fruit_scores) if fruit_scores else 0
        avg_cell = ws_fruits.cell(row=fruit_idx, column=avg_col)
        avg_cell.value = avg
        avg_cell.font = Font(bold=True)
        
        if avg > 0.05:
            avg_cell.fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
        elif avg < -0.05:
            avg_cell.fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
    
    ws_fruits.column_dimensions['A'].width = 15
    
    # ========================================================================
    # SHEET 3: Triad Components Detail
    # ========================================================================
    ws_triad = wb.create_sheet("Triad Components")
    
    ws_triad['A1'] = "Triad Component Breakdown (Pi, A, Lambda)"
    ws_triad['A1'].font = Font(size=14, bold=True, color="FFFFFF")
    ws_triad['A1'].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    ws_triad.merge_cells('A1:F1')
    
    row = 3
    for title, data in sorted(all_scores.items()):
        # Paper title
        ws_triad.cell(row=row, column=1, value=title)
        ws_triad.cell(row=row, column=1).font = Font(bold=True, size=12)
        row += 1
        
        # Pi components
        ws_triad.cell(row=row, column=2, value="Π (Polis):")
        ws_triad.cell(row=row, column=3, value=data['triad']['pi'])
        row += 1
        
        for comp_name, comp_val in data['triad']['pi_components'].items():
            ws_triad.cell(row=row, column=3, value=comp_name)
            ws_triad.cell(row=row, column=4, value=comp_val)
            row += 1
        
        # A components
        ws_triad.cell(row=row, column=2, value="A (Anthropos):")
        ws_triad.cell(row=row, column=3, value=data['triad']['a'])
        row += 1
        
        for comp_name, comp_val in data['triad']['a_components'].items():
            ws_triad.cell(row=row, column=3, value=comp_name)
            ws_triad.cell(row=row, column=4, value=comp_val)
            row += 1
        
        # Lambda components
        ws_triad.cell(row=row, column=2, value="Λ (Logos):")
        ws_triad.cell(row=row, column=3, value=data['triad']['lambda'])
        row += 1
        
        for comp_name, comp_val in data['triad']['lambda_components'].items():
            ws_triad.cell(row=row, column=3, value=comp_name)
            ws_triad.cell(row=row, column=4, value=comp_val)
            row += 1
        
        row += 1  # Spacer
    
    # ========================================================================
    # SHEET 4: Constraints
    # ========================================================================
    ws_constraints = wb.create_sheet("Constraints")
    
    ws_constraints['A1'] = "9 Structural Constraints - Satisfaction Matrix"
    ws_constraints['A1'].font = Font(size=14, bold=True, color="FFFFFF")
    ws_constraints['A1'].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    ws_constraints.merge_cells('A1:H1')
    
    # Get constraint names from first paper
    first_paper = list(all_scores.values())[0]
    constraint_names = [c['name'] for c in first_paper['constraints']]
    
    # Headers
    ws_constraints['A3'] = "Constraint"
    for col_idx, title in enumerate(sorted(all_scores.keys()), 2):
        cell = ws_constraints.cell(row=3, column=col_idx)
        cell.value = title[:20]
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    
    # Constraint rows
    for c_idx, c_name in enumerate(constraint_names, 4):
        ws_constraints.cell(row=c_idx, column=1, value=c_name)
        
        for col_idx, (title, data) in enumerate(sorted(all_scores.items()), 2):
            # Find constraint score
            for constraint in data['constraints']:
                if constraint['name'] == c_name:
                    cell = ws_constraints.cell(row=c_idx, column=col_idx)
                    score = constraint['score']
                    
                    if score == 1:
                        cell.value = "✓ Satisfied"
                        cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                    elif score == -1:
                        cell.value = "✗ Violated"
                        cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                        cell.font = Font(color="FF0000", bold=True)
                    else:
                        cell.value = "- Neutral"
                    break
    
    ws_constraints.column_dimensions['A'].width = 25
    
    # Save workbook
    output_path = outputs_dir / "Adam_Eve_Papers_Analysis.xlsx"
    wb.save(output_path)
    
    print(f"Excel report created: {output_path}")
    return output_path


if __name__ == "__main__":
    create_excel_report()
