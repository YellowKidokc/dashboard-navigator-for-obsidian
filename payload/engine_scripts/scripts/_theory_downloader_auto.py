#!/usr/bin/env python3
"""
Theory Downloader (AUTO MODE - no confirmation prompt)
Reads Excel file of consciousness theories, downloads each URL,
converts to clean Markdown, saves to output directory.

Usage:
  python O:\_Theophysics_v3\_theory_downloader_auto.py
"""
import json
import os
import re
import sys
import time
import subprocess

# ── Configuration ──────────────────────────────────────────────────────────
EXCEL_PATH = r"C:\Users\lowes\Downloads\landscape_consciousness_theories_links.xlsx"
OUTPUT_DIR = r"O:\999_IGNORE\Obsidian Programs\Theory_Downloader"
JSON_DUMP  = r"O:\_Theophysics_v3\_excel_output.json"
LOG_FILE   = r"O:\_Theophysics_v3\_download_log.txt"

# ── Phase 0: Install dependencies ─────────────────────────────────────────
def install_deps():
    required = {
        "openpyxl": "openpyxl",
        "requests": "requests",
        "bs4": "beautifulsoup4",
        "html2text": "html2text",
    }
    for import_name, pip_name in required.items():
        try:
            __import__(import_name)
        except ImportError:
            print(f"  Installing {pip_name}...")
            subprocess.check_call([sys.executable, "-m", "pip", "install", pip_name])
            print(f"  Installed {pip_name}")

print("Checking dependencies...")
install_deps()

import openpyxl
import requests
import html2text
from bs4 import BeautifulSoup

# ── Utilities ─────────────────────────────────────────────────────────────
def slugify(text, max_len=80):
    if not text:
        return "unnamed"
    text = re.sub(r'[^\w\s-]', '', str(text))
    text = re.sub(r'[-\s]+', '_', text).strip('_')
    return text[:max_len] or "unnamed"

def slug_from_url(url):
    from urllib.parse import urlparse
    parsed = urlparse(url)
    path = parsed.path.strip("/")
    if path:
        parts = [p for p in path.split("/") if p and p not in ("index.html", "index.htm", "index.php")]
        if parts:
            return slugify(parts[-1])
    return slugify(parsed.netloc.replace("www.", ""))

def download_and_convert(url, timeout=30):
    hdrs = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                       "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
    }
    resp = requests.get(url, headers=hdrs, timeout=timeout, allow_redirects=True)
    resp.raise_for_status()

    soup = BeautifulSoup(resp.text, "html.parser")

    # Remove junk
    for tag in soup.find_all(["script", "style", "nav", "footer", "header",
                               "aside", "iframe", "noscript", "form",
                               "button", "input", "select", "textarea"]):
        tag.decompose()

    for pattern in ["nav", "menu", "sidebar", "footer", "header", "ad-", "ads-",
                     "cookie", "popup", "modal", "banner", "social", "share",
                     "comment", "related"]:
        for el in soup.find_all(class_=re.compile(pattern, re.I)):
            el.decompose()
        for el in soup.find_all(id=re.compile(pattern, re.I)):
            el.decompose()

    # Find main content
    main = (soup.find("main") or soup.find("article") or
            soup.find("div", class_=re.compile(r"content|entry|post|article|body", re.I)) or
            soup.find("div", id=re.compile(r"content|entry|post|article|body", re.I)) or
            soup.body or soup)

    title = ""
    title_tag = soup.find("title")
    if title_tag:
        title = title_tag.get_text(strip=True)
    h1 = soup.find("h1")
    if h1:
        title = h1.get_text(strip=True)

    h = html2text.HTML2Text()
    h.ignore_links = False
    h.ignore_images = True
    h.ignore_emphasis = False
    h.body_width = 0
    h.skip_internal_links = True
    h.inline_links = True
    h.protect_links = True
    h.wrap_links = False
    h.wrap_list_items = False
    h.unicode_snob = True

    markdown = h.handle(str(main))
    markdown = re.sub(r'\n{4,}', '\n\n\n', markdown).strip()
    return title, markdown

# ── Phase 1: Read Excel ──────────────────────────────────────────────────
print(f"\n{'='*70}")
print(f"PHASE 1: Reading Excel")
print(f"{'='*70}")
print(f"  File: {EXCEL_PATH}")

wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb.active
print(f"  Sheet: {ws.title}, Rows: {ws.max_row}, Cols: {ws.max_column}")

headers = [cell.value for cell in ws[1]]
print(f"  Headers: {headers}")

rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    row_data = {}
    for i, val in enumerate(row):
        if i < len(headers):
            key = str(headers[i]) if headers[i] else f"col_{i}"
            row_data[key] = val
    if any(v is not None for v in row_data.values()):
        rows.append(row_data)

print(f"  Data rows: {len(rows)}")
print(f"\n{'─'*70}")
print("ALL ROWS:")
print(f"{'─'*70}")
for i, r in enumerate(rows, 1):
    print(f"\n  Row {i}:")
    for k, v in r.items():
        if v is not None:
            print(f"    {k}: {v}")

# Dump JSON
with open(JSON_DUMP, "w", encoding="utf-8") as f:
    json.dump({"headers": headers, "total_rows": len(rows), "rows": rows},
              f, indent=2, default=str, ensure_ascii=False)
print(f"\n  JSON: {JSON_DUMP}")

# ── Phase 2: Extract URLs ────────────────────────────────────────────────
print(f"\n{'='*70}")
print(f"PHASE 2: Extracting URLs")
print(f"{'='*70}")

url_cols = []
name_cols = []
cat_cols = []

for h in headers:
    if h is None:
        continue
    hl = str(h).lower()
    if any(kw in hl for kw in ["url", "link", "href", "website", "source", "reference"]):
        url_cols.append(str(h))
    if any(kw in hl for kw in ["name", "theory", "title", "label"]):
        name_cols.append(str(h))
    if any(kw in hl for kw in ["category", "type", "class", "group", "domain"]):
        cat_cols.append(str(h))

print(f"  URL cols: {url_cols}")
print(f"  Name cols: {name_cols}")
print(f"  Category cols: {cat_cols}")

entries = []
if url_cols:
    for row in rows:
        for uc in url_cols:
            url = row.get(uc)
            if url and str(url).startswith("http"):
                name = next((str(row[nc]) for nc in name_cols if row.get(nc)), None)
                cat = next((str(row[cc]) for cc in cat_cols if row.get(cc)), None)
                entries.append({
                    "url": str(url).strip(),
                    "name": name,
                    "category": cat,
                    "row_data": {k: str(v) for k, v in row.items() if v is not None}
                })
else:
    print("  No URL column found. Scanning all cells for URLs...")
    for row in rows:
        for h_key, v in row.items():
            if v and str(v).startswith("http"):
                name = next((str(v2) for h2, v2 in row.items()
                            if h2 != h_key and v2 and not str(v2).startswith("http")), None)
                entries.append({
                    "url": str(v).strip(),
                    "name": name,
                    "category": None,
                    "row_data": {k: str(val) for k, val in row.items() if val is not None}
                })

# Deduplicate
seen = set()
unique = []
for e in entries:
    if e["url"] not in seen:
        seen.add(e["url"])
        unique.append(e)
entries = unique

print(f"  Unique URLs: {len(entries)}")
for i, e in enumerate(entries, 1):
    print(f"  [{i}] {e['name'] or '(unnamed)'}: {e['url']}")

if not entries:
    print("\nNo URLs found! Check the Excel structure.")
    sys.exit(1)

# ── Phase 3: Download ────────────────────────────────────────────────────
print(f"\n{'='*70}")
print(f"PHASE 3: Downloading {len(entries)} URLs")
print(f"{'='*70}")

os.makedirs(OUTPUT_DIR, exist_ok=True)
print(f"  Output: {OUTPUT_DIR}\n")

log_lines = []
success = failed = skipped = 0

for i, entry in enumerate(entries, 1):
    url = entry["url"]
    name = entry["name"]
    category = entry.get("category")

    # Build filename
    if name:
        fname = slugify(name)
    else:
        fname = slug_from_url(url)
    if category:
        fname = slugify(category) + "__" + fname

    fpath = os.path.join(OUTPUT_DIR, fname + ".md")
    counter = 1
    while os.path.exists(fpath):
        fpath = os.path.join(OUTPUT_DIR, f"{fname}_{counter}.md")
        counter += 1

    print(f"  [{i}/{len(entries)}] {name or fname}")
    print(f"    URL: {url}")

    # Skip non-HTML
    if any(url.lower().endswith(ext) for ext in [".pdf", ".doc", ".docx", ".ppt", ".pptx", ".zip"]):
        msg = f"SKIPPED (non-HTML): {url}"
        print(f"    {msg}")
        log_lines.append(msg)
        skipped += 1
        with open(fpath, "w", encoding="utf-8") as f:
            f.write(f"# {name or fname}\n\n**Source:** {url}\n\n")
            if category:
                f.write(f"**Category:** {category}\n\n")
            f.write(f"*Direct download link. Visit URL to access.*\n")
        continue

    try:
        title, markdown = download_and_convert(url)

        md_content = f"# {title or name or fname}\n\n"
        md_content += f"**Source:** [{url}]({url})\n"
        if name:
            md_content += f"**Theory:** {name}\n"
        if category:
            md_content += f"**Category:** {category}\n"
        row_data = entry.get("row_data", {})
        extra = {k: v for k, v in row_data.items() if "url" not in k.lower() and "link" not in k.lower()}
        if extra:
            md_content += "\n**Metadata:**\n"
            for k, v in extra.items():
                md_content += f"- {k}: {v}\n"
        md_content += f"\n---\n\n{markdown}\n"

        with open(fpath, "w", encoding="utf-8") as f:
            f.write(md_content)

        sz = len(markdown)
        msg = f"OK ({sz:,} chars): {name or fname} -> {os.path.basename(fpath)}"
        print(f"    {msg}")
        log_lines.append(msg)
        success += 1
        time.sleep(1)

    except requests.exceptions.Timeout:
        msg = f"TIMEOUT: {url}"
        print(f"    {msg}")
        log_lines.append(msg)
        failed += 1
        with open(fpath, "w", encoding="utf-8") as f:
            f.write(f"# {name or fname}\n\n**Source:** {url}\n\n*Timed out.*\n")

    except requests.exceptions.HTTPError as e:
        code = e.response.status_code if e.response else "?"
        msg = f"HTTP {code}: {url}"
        print(f"    {msg}")
        log_lines.append(msg)
        failed += 1
        with open(fpath, "w", encoding="utf-8") as f:
            f.write(f"# {name or fname}\n\n**Source:** {url}\n\n*HTTP Error {code}.*\n")

    except Exception as e:
        msg = f"ERROR ({type(e).__name__}: {e}): {url}"
        print(f"    {msg}")
        log_lines.append(msg)
        failed += 1
        with open(fpath, "w", encoding="utf-8") as f:
            f.write(f"# {name or fname}\n\n**Source:** {url}\n\n*Error: {e}*\n")

# ── Summary ───────────────────────────────────────────────────────────────
print(f"\n{'='*70}")
print(f"SUMMARY")
print(f"{'='*70}")
print(f"  Total:      {len(entries)}")
print(f"  Success:    {success}")
print(f"  Failed:     {failed}")
print(f"  Skipped:    {skipped}")
print(f"  Output:     {OUTPUT_DIR}")

with open(LOG_FILE, "w", encoding="utf-8") as f:
    f.write(f"Theory Downloader Log - {time.strftime('%Y-%m-%d %H:%M:%S')}\n")
    f.write(f"Total: {len(entries)} | Success: {success} | Failed: {failed} | Skipped: {skipped}\n")
    f.write(f"{'='*70}\n")
    for line in log_lines:
        f.write(line + "\n")

print(f"  Log:        {LOG_FILE}")
print(f"\nDONE. Created {success + failed + skipped} .md files in Theory_Downloader/")
