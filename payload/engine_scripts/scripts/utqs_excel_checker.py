#!/usr/bin/env python3
"""
UTQS Excel checker.

Builds and scores a workbook for:
- Layer 1: Academic Standards (AS)
- Layer 2: Content Quality (TQD)
- Layer 3: System Quality (SQI)

Outputs:
- Updated workbook output sheets
- JSON report
- CSV remediation list
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import math
import os
import urllib.error
import urllib.request
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import Workbook, load_workbook


DOMAIN_DEFS = [
    {
        "code": "SI",
        "name": "Structural_Integrity",
        "weight": 0.25,
        "submetrics": [
            "Logic_Chain_Completeness",
            "Circular_Reasoning_Check",
            "Contradiction_Detection",
            "Premise_Clarity",
        ],
    },
    {
        "code": "EG",
        "name": "Empirical_Grounding",
        "weight": 0.20,
        "submetrics": [
            "Testability_Score",
            "Data_Quality",
            "Prediction_Precision",
            "Historical_Fit",
        ],
    },
    {
        "code": "FC",
        "name": "Framework_Coherence",
        "weight": 0.20,
        "submetrics": [
            "Master_Equation_Integration",
            "Ten_Laws_Alignment",
            "LLC_Connection",
            "Variable_Consistency",
        ],
    },
    {
        "code": "FL",
        "name": "Falsification_Clarity",
        "weight": 0.15,
        "submetrics": [
            "If_Correct_Specificity",
            "If_Wrong_Precision",
            "Timeline_Bounds",
            "Alternative_Explanations",
        ],
    },
    {
        "code": "TS",
        "name": "Theological_Soundness",
        "weight": 0.10,
        "submetrics": [
            "Scripture_Anchoring",
            "Doctrinal_Accuracy",
            "Boundary_Respect",
            "Tradition_Engagement",
        ],
    },
    {
        "code": "PR",
        "name": "Physical_Rigor",
        "weight": 0.10,
        "submetrics": [
            "Equation_Validity",
            "Analogy_vs_Isomorphism",
            "Mechanism_Specification",
            "Known_Physics_Consistency",
        ],
    },
]

SQI_DEFS = [
    ("Internal_Coherence", 0.25),
    ("Trace_Completeness", 0.20),
    ("Definition_Stability", 0.20),
    ("Predictive_Power", 0.20),
    ("Compression_Efficiency", 0.15),
]


DEFAULT_CONFIG = {
    "AS_GATE": 85,
    "TQD_DEV_GATE": 75,
    "TQD_PUB_GATE": 90,
    "SQI_GATE": 8.0,
    "MIN_DOMAIN_FLOOR": 70,
    "FC_FLOOR": 80,
    "FL_FLOOR": 80,
    "HARD_FAIL_CAP": 69,
    "REMEDIATION_TARGET": 90,
    "REMEDIATION_THRESHOLD": 85,
    "EVIDENCE_WEIGHT": 0.5,
    "AGREEMENT_WEIGHT": 0.5,
}


DEFAULT_FLAGS = [
    ("MISSING_FALSIFICATION_PATH", "CRITICAL", 0, "FL", "No explicit fatal falsification condition", 69),
    ("CIRCULAR_CORE_CLAIM", "CRITICAL", 0, "SI", "Core argument depends on its own conclusion", 69),
    ("FRAMEWORK_CONTRADICTION", "CRITICAL", 0, "FC", "Contradicts core framework constraints", 69),
    ("NO_REFERENCES_SECTION", "HIGH", 0, "EG", "No references section detected", 79),
    ("DOCTRINAL_ERROR_MAJOR", "HIGH", 0, "TS", "Major theological error detected", 79),
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="UTQS Excel scorer")
    parser.add_argument(
        "--workbook",
        default=r"O:\_Theophysics_v3\00_SYSTEM\00System\08_Quality_Scoring_Model\UTQS_Rating_Template.xlsx",
        help="Workbook path",
    )
    parser.add_argument("--init-workbook", action="store_true", help="Create template workbook")
    parser.add_argument("--score-workbook", action="store_true", help="Score workbook and write outputs")
    parser.add_argument("--openai-review", action="store_true", help="Run optional OpenAI adversarial review")
    parser.add_argument("--openai-model", default="gpt-4o-mini", help="OpenAI model for optional review")
    parser.add_argument(
        "--framework-context",
        default="",
        help="Optional framework context text file path for OpenAI review",
    )
    return parser.parse_args()


def domain_map() -> Dict[str, dict]:
    return {d["code"]: d for d in DOMAIN_DEFS}


def clamp_0_100(value: float) -> float:
    return max(0.0, min(100.0, value))


def to_float(value, default=0.0) -> float:
    if value is None:
        return float(default)
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return float(default)
    try:
        return float(s)
    except ValueError:
        return float(default)


def to_bool01(value) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    s = str(value).strip().lower()
    return s in {"1", "true", "yes", "y"}


def normalize_confidence_input(x: float) -> float:
    if x > 1.0:
        x = x / 100.0
    return max(0.0, min(1.0, x))


def score_band(score: float) -> str:
    if score >= 90:
        return "Publication-ready (Logos Papers tier)"
    if score >= 75:
        return "Strong, needs minor refinement"
    if score >= 60:
        return "Solid concept, structural gaps"
    if score >= 40:
        return "Salvageable with major rework"
    return "Fundamental problems, consider discarding"


def confidence_band(score: float) -> str:
    if score >= 85:
        return "HIGH"
    if score >= 65:
        return "MEDIUM"
    return "LOW"


def ensure_sheet(wb: Workbook, name: str):
    if name in wb.sheetnames:
        return wb[name]
    return wb.create_sheet(name)


def seed_empty_sheet_header(ws, headers: Sequence[str]) -> None:
    if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
        for idx, h in enumerate(headers, start=1):
            ws.cell(row=1, column=idx, value=h)


def clear_sheet_keep_header(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def write_rows(ws, rows: Iterable[Iterable]):
    for row in rows:
        ws.append(list(row))


def autosize(ws):
    widths: Dict[int, int] = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            val = "" if cell.value is None else str(cell.value)
            widths[cell.column] = max(widths.get(cell.column, 0), len(val))
    for col_idx, width in widths.items():
        ws.column_dimensions[chr(64 + col_idx)].width = min(max(width + 2, 10), 80)


def init_workbook(path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    ws_readme = wb.create_sheet("README")
    ws_readme.append(["UTQS Workbook"])
    ws_readme.append(
        ["1) Fill INPUT_PAPER, INPUT_SCORES, INPUT_FLAGS. 2) Run --score-workbook. 3) Review output sheets."]
    )
    ws_readme.append(["Hard fail flags can cap final score even if weighted average is high."])
    ws_readme.append(["Use SQI 0-10. Use TQD/AS 0-100."])

    ws_cfg = wb.create_sheet("CONFIG")
    ws_cfg.append(["parameter", "value"])
    for k, v in DEFAULT_CONFIG.items():
        ws_cfg.append([k, v])

    ws_dw = wb.create_sheet("DOMAIN_WEIGHTS")
    ws_dw.append(["domain_code", "domain_name", "weight"])
    for d in DOMAIN_DEFS:
        ws_dw.append([d["code"], d["name"], d["weight"]])

    ws_sub = wb.create_sheet("SUBMETRICS")
    ws_sub.append(["domain_code", "submetric", "submetric_weight"])
    for d in DOMAIN_DEFS:
        for sm in d["submetrics"]:
            ws_sub.append([d["code"], sm, 1.0])

    ws_p = wb.create_sheet("INPUT_PAPER")
    ws_p.append(
        [
            "paper_id",
            "paper_title",
            "paper_path",
            "author",
            "date",
            "paper_type",
            "as_score",
            "sqi_score",
            "evidence_coverage",
            "evaluator_agreement",
            "notes",
        ]
    )
    ws_p.append(
        [
            "SEM_01",
            "Semantic Entropy",
            "",
            "David Lowe",
            dt.date.today().isoformat(),
            "paper",
            85,
            8.0,
            0.7,
            0.8,
            "",
        ]
    )

    ws_sqi = wb.create_sheet("INPUT_SQI")
    ws_sqi.append(["dimension", "weight", "score_0_10", "notes"])
    for dim, w in SQI_DEFS:
        ws_sqi.append([dim, w, "", ""])

    ws_s = wb.create_sheet("INPUT_SCORES")
    ws_s.append(["domain_code", "submetric", "score_0_100", "line_refs", "notes"])
    for d in DOMAIN_DEFS:
        for sm in d["submetrics"]:
            ws_s.append([d["code"], sm, "", "", ""])

    ws_f = wb.create_sheet("INPUT_FLAGS")
    ws_f.append(["flag_code", "severity", "is_triggered_0_1", "domain_code", "description", "cap_score"])
    for row in DEFAULT_FLAGS:
        ws_f.append(list(row))

    ws_od = wb.create_sheet("OUTPUT_DOMAIN")
    ws_od.append(
        ["domain_code", "domain_name", "raw_domain_score", "weighted_contribution", "floor_threshold", "meets_floor"]
    )

    ws_of = wb.create_sheet("OUTPUT_FINAL")
    ws_of.append(["metric", "value"])

    ws_or = wb.create_sheet("OUTPUT_REMEDIATION")
    ws_or.append(
        [
            "priority",
            "type",
            "domain_code",
            "item",
            "current_score",
            "target_score",
            "estimated_impact",
            "action",
        ]
    )

    ws_oa = wb.create_sheet("OUTPUT_AI_REVIEW")
    ws_oa.append(["section", "content"])
    ws_oa.append(["status", "Not run"])

    ws_osqi = wb.create_sheet("OUTPUT_SQI")
    ws_osqi.append(["dimension", "weight", "score", "weighted_contribution"])

    for name in wb.sheetnames:
        autosize(wb[name])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_config(ws_cfg) -> Dict[str, float]:
    cfg = dict(DEFAULT_CONFIG)
    for row in ws_cfg.iter_rows(min_row=2, values_only=True):
        key, value = row[:2]
        if not key:
            continue
        cfg[str(key).strip()] = to_float(value, cfg.get(str(key).strip(), 0.0))
    return cfg


def read_domain_weights(ws_dw) -> Dict[str, float]:
    weights = {}
    for row in ws_dw.iter_rows(min_row=2, values_only=True):
        code, _, weight = row[:3]
        if not code:
            continue
        weights[str(code).strip()] = to_float(weight, 0.0)
    return weights


def read_submetric_weights(ws_sub) -> Dict[Tuple[str, str], float]:
    out: Dict[Tuple[str, str], float] = {}
    for row in ws_sub.iter_rows(min_row=2, values_only=True):
        code, sm, w = row[:3]
        if not code or not sm:
            continue
        out[(str(code).strip(), str(sm).strip())] = to_float(w, 1.0)
    return out


def read_input_paper(ws_p) -> Dict[str, str]:
    headers = [str(c.value).strip() if c.value else "" for c in ws_p[1]]
    values = [c.value for c in ws_p[2]]
    out = {}
    for i, h in enumerate(headers):
        if h:
            out[h] = values[i] if i < len(values) else None
    return out


def read_input_scores(ws_s) -> List[dict]:
    rows = []
    for row in ws_s.iter_rows(min_row=2, values_only=True):
        code, submetric, score, line_refs, notes = row[:5]
        if not code or not submetric:
            continue
        rows.append(
            {
                "domain_code": str(code).strip(),
                "submetric": str(submetric).strip(),
                "score": to_float(score, 0.0),
                "line_refs": "" if line_refs is None else str(line_refs),
                "notes": "" if notes is None else str(notes),
            }
        )
    return rows


def read_input_flags(ws_f) -> List[dict]:
    rows = []
    for row in ws_f.iter_rows(min_row=2, values_only=True):
        code, severity, triggered, domain_code, desc, cap = row[:6]
        if not code:
            continue
        rows.append(
            {
                "flag_code": str(code).strip(),
                "severity": str(severity or "").strip().upper(),
                "triggered": to_bool01(triggered),
                "domain_code": str(domain_code or "").strip(),
                "description": str(desc or "").strip(),
                "cap_score": to_float(cap, math.nan),
            }
        )
    return rows


def read_input_sqi(ws_sqi) -> List[dict]:
    rows = []
    for row in ws_sqi.iter_rows(min_row=2, values_only=True):
        dim, weight, score, notes = row[:4]
        if not dim:
            continue
        rows.append(
            {
                "dimension": str(dim).strip(),
                "weight": to_float(weight, 0.0),
                "score": to_float(score, math.nan),
                "notes": str(notes or "").strip(),
            }
        )
    return rows


def weighted_average(items: List[Tuple[float, float]]) -> float:
    total_w = sum(w for _, w in items)
    if total_w <= 0:
        return 0.0
    return sum(v * w for v, w in items) / total_w


def run_scoring(workbook_path: Path, openai_review: bool, openai_model: str, framework_context_file: str) -> Dict:
    wb = load_workbook(workbook_path)
    ws_cfg = wb["CONFIG"]
    ws_dw = wb["DOMAIN_WEIGHTS"]
    ws_sub = wb["SUBMETRICS"]
    ws_p = wb["INPUT_PAPER"]
    ws_s = wb["INPUT_SCORES"]
    ws_f = wb["INPUT_FLAGS"]
    ws_sqi = ensure_sheet(wb, "INPUT_SQI")
    ws_od = wb["OUTPUT_DOMAIN"]
    ws_of = wb["OUTPUT_FINAL"]
    ws_or = wb["OUTPUT_REMEDIATION"]
    ws_oa = wb["OUTPUT_AI_REVIEW"]
    ws_osqi = ensure_sheet(wb, "OUTPUT_SQI")

    seed_empty_sheet_header(ws_sqi, ["dimension", "weight", "score_0_10", "notes"])
    if ws_sqi.max_row == 1:
        for dim, w in SQI_DEFS:
            ws_sqi.append([dim, w, "", ""])
    seed_empty_sheet_header(ws_osqi, ["dimension", "weight", "score", "weighted_contribution"])

    cfg = read_config(ws_cfg)
    d_weights = read_domain_weights(ws_dw)
    sm_weights = read_submetric_weights(ws_sub)
    paper = read_input_paper(ws_p)
    score_rows = read_input_scores(ws_s)
    flags = read_input_flags(ws_f)
    sqi_rows = read_input_sqi(ws_sqi)

    dmap = domain_map()
    by_domain: Dict[str, List[Tuple[float, float, str]]] = {d["code"]: [] for d in DOMAIN_DEFS}
    for row in score_rows:
        code = row["domain_code"]
        if code not in by_domain:
            continue
        key = (code, row["submetric"])
        w = sm_weights.get(key, 1.0)
        by_domain[code].append((clamp_0_100(row["score"]), w, row["submetric"]))

    domain_scores: Dict[str, float] = {}
    domain_contrib: Dict[str, float] = {}
    for code, items in by_domain.items():
        if items:
            avg = weighted_average([(v, w) for v, w, _ in items])
        else:
            avg = 0.0
        domain_scores[code] = avg
        domain_contrib[code] = avg * d_weights.get(code, dmap.get(code, {}).get("weight", 0.0))

    tqd_raw = sum(domain_contrib.values())
    tqd_final = tqd_raw

    critical_flags = [f for f in flags if f["triggered"] and f["severity"] in {"CRITICAL", "BLOCKER"}]
    cap_applied = None
    if critical_flags:
        cap_candidates = [cfg.get("HARD_FAIL_CAP", 69.0)]
        for f in critical_flags:
            if not math.isnan(f["cap_score"]):
                cap_candidates.append(f["cap_score"])
        cap_applied = min(cap_candidates)
        tqd_final = min(tqd_final, cap_applied)

    as_score = clamp_0_100(to_float(paper.get("as_score"), 0.0))

    sqi_components = []
    for r in sqi_rows:
        if math.isnan(r["score"]):
            continue
        sc = max(0.0, min(10.0, r["score"]))
        wt = r["weight"] if r["weight"] > 0 else 0.0
        sqi_components.append((r["dimension"], wt, sc))

    if sqi_components:
        sqi_score = weighted_average([(sc, wt) for _, wt, sc in sqi_components])
    else:
        sqi_score = max(0.0, min(10.0, to_float(paper.get("sqi_score"), 0.0)))

    min_floor = cfg.get("MIN_DOMAIN_FLOOR", 70.0)
    fc_floor = cfg.get("FC_FLOOR", 80.0)
    fl_floor = cfg.get("FL_FLOOR", 80.0)
    floor_failures = []
    for code, val in domain_scores.items():
        threshold = min_floor
        if code == "FC":
            threshold = fc_floor
        elif code == "FL":
            threshold = fl_floor
        if val < threshold:
            floor_failures.append((code, val, threshold))

    as_gate = cfg.get("AS_GATE", 85.0)
    dev_gate = cfg.get("TQD_DEV_GATE", 75.0)
    pub_gate = cfg.get("TQD_PUB_GATE", 90.0)
    sqi_gate = cfg.get("SQI_GATE", 8.0)

    if as_score < as_gate:
        status = "REJECTED_LAYER_1"
    elif tqd_final < dev_gate:
        status = "NEEDS_REVISION"
    elif tqd_final >= pub_gate and sqi_score >= sqi_gate and not floor_failures and not critical_flags:
        status = "PUBLICATION_READY"
    else:
        status = "STRONG_NEEDS_REFINEMENT"

    ev_cov = normalize_confidence_input(to_float(paper.get("evidence_coverage"), 0.7))
    ev_agree = normalize_confidence_input(to_float(paper.get("evaluator_agreement"), 0.7))
    conf_score = clamp_0_100(
        100
        * (
            cfg.get("EVIDENCE_WEIGHT", 0.5) * ev_cov
            + cfg.get("AGREEMENT_WEIGHT", 0.5) * ev_agree
        )
    )
    conf_level = confidence_band(conf_score)

    # Build remediation list
    remediation = []
    priority = 1
    for f in sorted(
        [x for x in flags if x["triggered"]],
        key=lambda x: {"CRITICAL": 0, "BLOCKER": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3}.get(x["severity"], 4),
    ):
        impact = 8 if f["severity"] in {"CRITICAL", "BLOCKER"} else (5 if f["severity"] == "HIGH" else 3)
        remediation.append(
            {
                "priority": priority,
                "type": "FLAG",
                "domain_code": f["domain_code"],
                "item": f["flag_code"],
                "current_score": "",
                "target_score": "",
                "estimated_impact": impact,
                "action": f["description"] or "Resolve flagged critical issue",
            }
        )
        priority += 1

    target = cfg.get("REMEDIATION_TARGET", 90.0)
    threshold = cfg.get("REMEDIATION_THRESHOLD", 85.0)
    for row in sorted(score_rows, key=lambda x: x["score"]):
        if row["score"] >= threshold:
            continue
        code = row["domain_code"]
        weight = d_weights.get(code, dmap.get(code, {}).get("weight", 0.0))
        n_sub = max(1, len(dmap.get(code, {}).get("submetrics", [])))
        needed = max(0.0, target - row["score"])
        est_impact = round((needed * weight / n_sub), 2)
        remediation.append(
            {
                "priority": priority,
                "type": "SUBMETRIC",
                "domain_code": code,
                "item": row["submetric"],
                "current_score": round(row["score"], 2),
                "target_score": target,
                "estimated_impact": est_impact,
                "action": f"Raise {row['submetric']} to >= {target}.",
            }
        )
        priority += 1

    # Optional OpenAI review
    ai_review = {"status": "Not run"}
    if openai_review:
        ai_review = run_openai_review(
            paper_path=str(paper.get("paper_path") or ""),
            framework_context_file=framework_context_file,
            model=openai_model,
        )

    # Write outputs
    clear_sheet_keep_header(ws_od)
    for code in [d["code"] for d in DOMAIN_DEFS]:
        d = dmap[code]
        threshold = min_floor
        if code == "FC":
            threshold = fc_floor
        elif code == "FL":
            threshold = fl_floor
        ws_od.append(
            [
                code,
                d["name"],
                round(domain_scores.get(code, 0.0), 2),
                round(domain_contrib.get(code, 0.0), 2),
                threshold,
                "YES" if domain_scores.get(code, 0.0) >= threshold else "NO",
            ]
        )

    clear_sheet_keep_header(ws_of)
    out_rows = [
        ("paper_id", paper.get("paper_id", "")),
        ("paper_title", paper.get("paper_title", "")),
        ("as_score", round(as_score, 2)),
        ("sqi_score", round(sqi_score, 2)),
        ("tqd_raw", round(tqd_raw, 2)),
        ("tqd_final", round(tqd_final, 2)),
        ("score_band", score_band(tqd_final)),
        ("status", status),
        ("domain_floor_pass", "YES" if not floor_failures else "NO"),
        (
            "domain_floor_failures",
            "; ".join([f"{c}:{round(v,1)}<{t}" for c, v, t in floor_failures]) if floor_failures else "",
        ),
        ("critical_fail_count", len(critical_flags)),
        ("cap_applied", "" if cap_applied is None else cap_applied),
        ("confidence_score", round(conf_score, 2)),
        ("confidence_level", conf_level),
        ("generated_at", dt.datetime.now().isoformat(timespec="seconds")),
    ]
    for r in out_rows:
        ws_of.append(list(r))

    clear_sheet_keep_header(ws_or)
    for r in remediation:
        ws_or.append(
            [
                r["priority"],
                r["type"],
                r["domain_code"],
                r["item"],
                r["current_score"],
                r["target_score"],
                r["estimated_impact"],
                r["action"],
            ]
        )

    clear_sheet_keep_header(ws_oa)
    for k, v in ai_review.items():
        if isinstance(v, (dict, list)):
            ws_oa.append([k, json.dumps(v, ensure_ascii=False)])
        else:
            ws_oa.append([k, str(v)])

    clear_sheet_keep_header(ws_osqi)
    if sqi_components:
        for dim, wt, sc in sqi_components:
            ws_osqi.append([dim, wt, round(sc, 2), round(sc * wt, 3)])
    else:
        ws_osqi.append(["SQI_MANUAL", 1.0, round(sqi_score, 2), round(sqi_score, 2)])

    for name in ["OUTPUT_DOMAIN", "OUTPUT_FINAL", "OUTPUT_REMEDIATION", "OUTPUT_AI_REVIEW"]:
        autosize(wb[name])
    autosize(ws_osqi)

    wb.save(workbook_path)

    report = {
        "paper_id": paper.get("paper_id", ""),
        "paper_title": paper.get("paper_title", ""),
        "as_score": round(as_score, 2),
        "sqi_score": round(sqi_score, 2),
        "tqd_raw": round(tqd_raw, 2),
        "tqd_final": round(tqd_final, 2),
        "status": status,
        "score_band": score_band(tqd_final),
        "domain_scores": {k: round(v, 2) for k, v in domain_scores.items()},
        "floor_failures": [
            {"domain_code": c, "score": round(v, 2), "threshold": t} for c, v, t in floor_failures
        ],
        "critical_flags": critical_flags,
        "confidence_score": round(conf_score, 2),
        "confidence_level": conf_level,
        "sqi_components": [
            {"dimension": dim, "weight": wt, "score": round(sc, 2)}
            for dim, wt, sc in sqi_components
        ],
        "remediation_count": len(remediation),
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
    }
    return report


def run_openai_review(paper_path: str, framework_context_file: str, model: str) -> Dict[str, str]:
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        return {"status": "Skipped", "reason": "OPENAI_API_KEY not set"}
    p = Path(paper_path) if paper_path else None
    if not p or not p.exists():
        return {"status": "Skipped", "reason": "paper_path missing or not found in INPUT_PAPER"}

    text = p.read_text(encoding="utf-8", errors="ignore")
    text = text[:22000]
    framework = ""
    if framework_context_file:
        f = Path(framework_context_file)
        if f.exists():
            framework = f.read_text(encoding="utf-8", errors="ignore")[:12000]

    prompt = (
        "Perform a strict adversarial review.\n"
        "Return JSON with keys: logical_flaws, empirical_weaknesses, "
        "framework_violations, theological_issues, blind_spots.\n\n"
        f"Framework Context:\n{framework}\n\n"
        f"Paper:\n{text}"
    )

    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": "You are a rigorous adversarial reviewer."},
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.2,
    }

    req = urllib.request.Request(
        "https://api.openai.com/v1/chat/completions",
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            body = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return {"status": "Error", "reason": f"HTTP {e.code}: {e.reason}"}
    except Exception as e:
        return {"status": "Error", "reason": str(e)}

    try:
        content = body["choices"][0]["message"]["content"]
    except Exception:
        return {"status": "Error", "reason": "Unexpected API response", "raw": body}

    return {"status": "OK", "model": model, "review": content}


def save_reports(workbook_path: Path, report: Dict) -> Tuple[Path, Path]:
    ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = workbook_path.parent
    json_path = out_dir / f"UTQS_REPORT_{ts}.json"
    csv_path = out_dir / f"UTQS_REMEDIATION_{ts}.csv"

    with json_path.open("w", encoding="utf-8") as fh:
        json.dump(report, fh, indent=2)

    # CSV remediation from workbook output sheet
    wb = load_workbook(workbook_path)
    ws = wb["OUTPUT_REMEDIATION"]
    with csv_path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["priority", "type", "domain_code", "item", "current_score", "target_score", "estimated_impact", "action"])
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(v is not None and str(v).strip() != "" for v in row):
                continue
            writer.writerow(list(row))
    return json_path, csv_path


def main() -> int:
    args = parse_args()
    workbook_path = Path(args.workbook)

    if not args.init_workbook and not args.score_workbook:
        print("No action selected. Use --init-workbook and/or --score-workbook.")
        return 1

    if args.init_workbook:
        init_workbook(workbook_path)
        print(f"Initialized workbook: {workbook_path}")

    if args.score_workbook:
        if not workbook_path.exists():
            print(f"Workbook not found: {workbook_path}")
            return 1
        report = run_scoring(
            workbook_path=workbook_path,
            openai_review=args.openai_review,
            openai_model=args.openai_model,
            framework_context_file=args.framework_context,
        )
        json_path, csv_path = save_reports(workbook_path, report)
        print(f"Scored workbook: {workbook_path}")
        print(f"Status: {report.get('status')} | TQD={report.get('tqd_final')} | AS={report.get('as_score')} | SQI={report.get('sqi_score')}")
        print(f"Report JSON: {json_path}")
        print(f"Remediation CSV: {csv_path}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
